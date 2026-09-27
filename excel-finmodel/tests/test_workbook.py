"""Структура книги: листы, имена, живые формулы без зашитых чисел, защита, проверки ввода."""
from pathlib import Path

import pytest
from openpyxl import load_workbook
from openpyxl.formula.tokenizer import Token, Tokenizer
from openpyxl.worksheet.formula import DataTableFormula

from conftest import simple_config
from finmodel import build as B
from finmodel.build import build_workbook
from finmodel.config import demo_config
from finmodel.mirror import compute

EXAMPLE = Path(__file__).resolve().parent.parent / "examples" / "finmodel-import-china.xlsx"
ALLOWED_LITERALS = {0.0, 1.0, 2.0, 3.0}   # ROUND(x,0), месяц−1, квартал = 3 мес.
CALC_SHEETS = [B.SH_DASH, B.SH_UNIT, B.SH_FLOW, B.SH_CASH, B.SH_PNL]
REQUIRED_NAMES = [
    "Дата_начала", "Капитал_старт", "Курс_база", "Курс", "Доставка_способ", "Налог_режим",
    "Темп_продаж", "Сценарий_выбор", "Сценарий_номер", "Множитель_цены", "Множитель_продаж",
    "ДРР", "Курс_анализ", "Цена_анализ", "Объём_анализ", "ЧП_итого", "Выручка_итого",
    "ROI_период", "Пик_вложений", "Макс_разрыв", "Окупаемость_мес", "Контроль_баланса",
]


def cells(ws):
    for row in ws.iter_rows():
        for c in row:
            if c.value is not None:
                yield c


def test_sheets_in_order(built):
    wb, _ = built
    assert wb.sheetnames == B.SHEET_ORDER
    assert wb.active.title == B.SH_DASH


def test_defined_names(built):
    wb, L = built
    for nm in REQUIRED_NAMES:
        assert nm in wb.defined_names, nm
        sheet, addr = L.names[nm]
        assert wb.defined_names[nm].attr_text == f"'{sheet}'!{addr}"
    assert L.names["Темп_продаж"][1].count(":") == 1          # диапазон на 18 месяцев
    assert L.names["ЧП_итого"][0] == B.SH_DASH


def test_no_hardcoded_numbers(built):
    """Числа-константы — только в ячейках ввода и служебных ячейках анализа."""
    wb, L = built
    bad = []
    for ws in wb.worksheets:
        for c in cells(ws):
            if isinstance(c.value, (int, float)) and not isinstance(c.value, bool):
                if (ws.title, c.coordinate) not in L.inputs | L.service:
                    bad.append(f"{ws.title}!{c.coordinate}={c.value}")
    assert not bad, bad[:20]


def test_formula_literals_whitelisted(built):
    wb, _ = built
    bad = []
    n = 0
    for ws in wb.worksheets:
        for c in cells(ws):
            if isinstance(c.value, str) and c.value.startswith("="):
                n += 1
                assert "#REF" not in c.value, c.coordinate
                for tok in Tokenizer(c.value).items:
                    if tok.type == Token.OPERAND and tok.subtype == Token.NUMBER:
                        if float(tok.value) not in ALLOWED_LITERALS:
                            bad.append(f"{ws.title}!{c.coordinate}: {tok.value} in {c.value}")
    assert n > 2000
    assert not bad, bad[:10]


def test_calc_sheets_are_formulas(built):
    """На расчётных листах каждое число получается формулой."""
    wb, L = built
    for name in CALC_SHEETS:
        ws = wb[name]
        n_formula = sum(1 for c in cells(ws) if isinstance(c.value, str) and c.value[:1] == "=")
        assert n_formula > 50, name
        for c in cells(ws):
            assert not isinstance(c.value, (int, float)), f"{name}!{c.coordinate}"


def test_protection_and_input_cells(built):
    wb, L = built
    for ws in wb.worksheets:
        assert ws.protection.sheet, ws.title
        assert not ws.protection.password, ws.title      # без пароля — снимается в один клик
    assert len(L.inputs) > 100
    for sheet, coord in L.inputs:
        c = wb[sheet][coord]
        assert c.protection.locked is False, coord
        assert c.fill.fgColor.rgb.endswith("FFF2CC"), coord
    # расчётные ячейки защищены
    for name in CALC_SHEETS:
        for c in cells(wb[name]):
            assert c.protection.locked, f"{name}!{c.coordinate}"
    # ввод — только на «Допущениях», «Сценариях» и осях «Чувствительности»
    assert {s for s, _ in L.inputs} == {B.SH_IN, B.SH_SC, B.SH_SENS}


def test_switches_have_dropdowns(built):
    wb, L = built
    for nm in ("Сценарий_выбор", "Налог_режим", "Доставка_способ"):
        sheet, addr = L.names[nm]
        coord = addr.replace("$", "")
        lists = [dv for dv in wb[sheet].data_validations.dataValidation
                 if dv.type == "list" and coord in dv.sqref]
        assert lists, nm


def test_cash_gap_highlight(built):
    wb, L = built
    row = L.cash_rows["closing"]
    rng = f"{B.FIRST_M}{row}:{B.LAST_M}{row}"
    rules = [r for cf in wb[B.SH_CASH].conditional_formatting if str(cf.sqref) == rng
             for r in cf.rules]
    assert any(r.operator == "lessThan" and r.formula == ["0"] for r in rules)


def test_charts_and_data_tables(built):
    wb, L = built
    assert len(wb[B.SH_DASH]._charts) == 3
    tables = [(ws.title, c.coordinate) for ws in wb.worksheets for c in cells(ws)
              if isinstance(c.value, DataTableFormula)]
    assert len(tables) == 4          # 3 таблицы чувствительности + сравнение сценариев
    assert {t[0] for t in tables} == {B.SH_SENS, B.SH_SC}


@pytest.mark.parametrize("n", [1, 3, 8])
def test_any_number_of_skus(n, tmp_path):
    base = demo_config()
    skus = (base.skus * 2)[:n]
    a = base.copy(skus=skus)
    wb, L = build_workbook(a)
    assert L.n_skus == n and len(L.flow_rows["orders"]) == n
    wb.save(tmp_path / "m.xlsx")
    load_workbook(tmp_path / "m.xlsx")      # файл читается
    compute(a)                              # и зеркало считается


def test_simple_config_builds(tmp_path):
    wb, _ = build_workbook(simple_config())
    wb.save(tmp_path / "simple.xlsx")


@pytest.mark.skipif(not EXAMPLE.exists(), reason="пример ещё не собран")
def test_example_file_has_values_from_excel():
    """Готовый пример сохранён Excel: в нём есть посчитанные значения, и они сходятся с зеркалом."""
    wb = load_workbook(EXAMPLE, data_only=True)
    r = compute(demo_config())
    _, L = build_workbook(demo_config())
    for key, nm in (("net_profit", "ЧП_итого"), ("revenue", "Выручка_итого"),
                    ("max_gap", "Макс_разрыв"), ("payback", "Окупаемость_мес")):
        sheet, addr = L.names[nm]
        assert wb[sheet][addr.replace("$", "")].value == pytest.approx(r["kpi"][key], abs=0.01)
    wbf = load_workbook(EXAMPLE)
    assert wbf.sheetnames == B.SHEET_ORDER
    for nm in REQUIRED_NAMES:
        assert nm in wbf.defined_names
    assert all(ws.protection.sheet for ws in wbf.worksheets)
