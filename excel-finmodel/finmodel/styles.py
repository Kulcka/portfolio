"""Оформление книги: цвета, шрифты, форматы чисел."""
from openpyxl.styles import Alignment, Border, Font, PatternFill, Protection, Side

FONT_NAME = "Calibri"

# палитра
C_NAVY = "1F3864"
C_BLUE = "2F5597"
C_INPUT_FILL = "FFF2CC"     # ввод — светло-жёлтый
C_INPUT_FONT = "1F4E79"     # ввод — тёмно-синий шрифт
C_SECTION = "D9E1F2"
C_TOTAL = "F2F2F2"
C_SERVICE = "EDEDED"
C_GRID = "D9D9D9"
C_RED = "C00000"
C_RED_FILL = "F8CBAD"
C_GREEN = "548235"
C_GREEN_FILL = "E2EFDA"
C_GRAY_TEXT = "7F7F7F"

thin = Side(style="thin", color=C_GRID)
BORDER = Border(left=thin, right=thin, top=thin, bottom=thin)
BORDER_BOTTOM = Border(bottom=Side(style="thin", color=C_BLUE))

# форматы чисел (хранятся в файле в «английской» записи, Excel показывает по локали)
F_RUB = '#,##0;[Red]-#,##0;"–"'
F_RUB_KPI = '#,##0" ₽";[Red]-#,##0" ₽";"0 ₽"'
F_RUB_UNIT = '#,##0.0;[Red]-#,##0.0;"–"'
F_QTY = '#,##0;[Red]-#,##0;"–"'
F_QTY1 = '#,##0.0;[Red]-#,##0.0;"–"'
F_PCT = '0.0%;[Red]-0.0%;"–"'
F_PCT_IN = '0.0%'
F_PCT_DELTA = '+0.0%;[Red]-0.0%;0.0%'
F_FX = '0.00'
F_CNY = '#,##0.00'
F_DATE = 'mmm yy'
F_DATE_FULL = 'dd.mm.yyyy'
F_DATE_MONTH = 'mmmm yyyy'
F_INT = '0'
F_M3 = '0.0000'
F_KG = '0.00'

FILL_INPUT = PatternFill("solid", fgColor=C_INPUT_FILL)
FILL_SECTION = PatternFill("solid", fgColor=C_SECTION)
FILL_TOTAL = PatternFill("solid", fgColor=C_TOTAL)
FILL_SERVICE = PatternFill("solid", fgColor=C_SERVICE)
FILL_HEADER = PatternFill("solid", fgColor=C_NAVY)
FILL_RED = PatternFill("solid", fgColor=C_RED_FILL)
FILL_GREEN = PatternFill("solid", fgColor=C_GREEN_FILL)
FILL_TILE = PatternFill("solid", fgColor="F3F6FB")

FONT = Font(name=FONT_NAME, size=10)
FONT_BOLD = Font(name=FONT_NAME, size=10, bold=True)
FONT_INPUT = Font(name=FONT_NAME, size=10, color=C_INPUT_FONT)
FONT_INPUT_BOLD = Font(name=FONT_NAME, size=10, color=C_INPUT_FONT, bold=True)
FONT_HEADER = Font(name=FONT_NAME, size=10, bold=True, color="FFFFFF")
FONT_TITLE = Font(name=FONT_NAME, size=16, bold=True, color=C_NAVY)
FONT_SUBTITLE = Font(name=FONT_NAME, size=11, bold=True, color=C_NAVY)
FONT_NOTE = Font(name=FONT_NAME, size=9, italic=True, color=C_GRAY_TEXT)
FONT_DEMO = Font(name=FONT_NAME, size=10, italic=True, bold=True, color=C_RED)
FONT_SERVICE = Font(name=FONT_NAME, size=9, italic=True, color=C_GRAY_TEXT)
FONT_KPI = Font(name=FONT_NAME, size=18, bold=True, color=C_NAVY)
FONT_KPI_LABEL = Font(name=FONT_NAME, size=9, bold=True, color=C_GRAY_TEXT)
FONT_LINK = Font(name=FONT_NAME, size=10, color="0563C1", underline="single")
FONT_WARN = Font(name=FONT_NAME, size=11, bold=True, color=C_RED)

UNLOCKED = Protection(locked=False)
LOCKED = Protection(locked=True)

AL_LEFT = Alignment(horizontal="left", vertical="center")
AL_WRAP = Alignment(horizontal="left", vertical="top", wrap_text=True)
AL_CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)
AL_RIGHT = Alignment(horizontal="right", vertical="center")
AL_INDENT = Alignment(horizontal="left", vertical="center", indent=2)
