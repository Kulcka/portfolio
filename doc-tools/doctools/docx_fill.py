"""docx-fill: массовое заполнение шаблона Word данными из таблицы.

В шаблоне — плейсхолдеры {{Название колонки}}, можно с форматом:
  {{Сумма|деньги}}    -> 1 250 000,00
  {{Сумма|прописью}}  -> Один миллион двести пятьдесят тысяч рублей 00 копеек
  {{Дата|дата}}       -> 01.10.2025
  {{Число|число}}     -> 1 250 000
  {{ФИО|инициалы}}    -> Иванова А. С.
  {{Поле|заглавные}}, {{Поле|строчные}}
Плейсхолдеры работают в тексте, таблицах, колонтитулах и надписях, даже если
Word разбил их на несколько фрагментов (после правок, проверки орфографии).
Форматирование берётся от первого символа плейсхолдера.
"""

from __future__ import annotations

import re
from dataclasses import dataclass, field
from datetime import date, datetime
from pathlib import Path
from typing import Any, Callable, Iterator

from docx import Document
from docx.oxml.ns import qn

from . import DocToolsError
from .normalize import (amount_in_words, format_date, format_money, format_number, initials, is_empty,
                        norm_key)
from .report import Report
from .tableio import read_table, to_text

PLACEHOLDER = re.compile(r"\{\{\s*([^{}|]+?)\s*(?:\|\s*([^{}]+?)\s*)?\}\}")
XML_SPACE = "{http://www.w3.org/XML/1998/namespace}space"

FORMATS: dict[str, Callable[[Any], str]] = {
    "деньги": format_money,
    "прописью": amount_in_words,
    "дата": format_date,
    "число": format_number,
    "инициалы": initials,
    "заглавные": lambda v: to_text(v).upper(),
    "строчные": lambda v: to_text(v).lower(),
}


def default_text(v: Any) -> str:
    if v is None:
        return ""
    if isinstance(v, (date, datetime)):
        return format_date(v)
    if isinstance(v, float):
        return format_number(v) if not v.is_integer() else format_number(int(v))
    if isinstance(v, int) and not isinstance(v, bool):
        return str(v)
    return str(v)


# ------------------------------------------------------------ обход документа
def _parts(doc: Any) -> Iterator[Any]:
    """Корневые XML-элементы: тело и все колонтитулы всех разделов."""
    yield doc.element.body
    seen = set()
    for sec in doc.sections:
        for hf in (sec.header, sec.first_page_header, sec.even_page_header,
                   sec.footer, sec.first_page_footer, sec.even_page_footer):
            if hf.is_linked_to_previous:
                continue
            el = hf._element
            if id(el) not in seen:
                seen.add(id(el))
                yield el


def _paragraphs(doc: Any) -> Iterator[Any]:
    for root in _parts(doc):
        yield from root.iter(qn("w:p"))


def _own_text_nodes(p: Any) -> list[Any]:
    """w:t этого абзаца (без текста вложенных надписей — у них свои абзацы)."""
    out = []
    for t in p.iter(qn("w:t")):
        parent = t.getparent()
        while parent is not None and parent.tag != qn("w:p"):
            parent = parent.getparent()
        if parent is p:
            out.append(t)
    return out


def find_placeholders(doc: Any) -> list[tuple[str, str | None]]:
    found: list[tuple[str, str | None]] = []
    for p in _paragraphs(doc):
        text = "".join(t.text or "" for t in _own_text_nodes(p))
        for m in PLACEHOLDER.finditer(text):
            item = (m.group(1), m.group(2))
            if item not in found:
                found.append(item)
    return found


def _replace_in_paragraph(p: Any, render: Callable[[str, str | None], str]) -> int:
    """Замена плейсхолдеров, в том числе разбитых по нескольким w:t (фрагментам текста)."""
    nodes = _own_text_nodes(p)
    if not nodes:
        return 0
    full = "".join(t.text or "" for t in nodes)
    matches = list(PLACEHOLDER.finditer(full))
    for m in reversed(matches):
        start, end = m.span()
        value = render(m.group(1), m.group(2))
        pos = 0
        bounds = []
        for t in nodes:
            n = len(t.text or "")
            bounds.append((pos, pos + n))
            pos += n
        si = next(i for i, (a, b) in enumerate(bounds) if a <= start < b)
        ei = next(i for i, (a, b) in enumerate(bounds) if a < end <= b)
        for i in range(si, ei + 1):
            a, _ = bounds[i]
            text = nodes[i].text or ""
            left = text[: start - a] if i == si else ""
            right = text[end - a:] if i == ei else ""
            nodes[i].text = left + (value if i == si else "") + right
            if nodes[i].text != nodes[i].text.strip():
                nodes[i].set(XML_SPACE, "preserve")
    return len(matches)


# ------------------------------------------------------------ заполнение
def _safe_filename(name: str) -> str:
    name = re.sub(r'[<>:"/\\|?*\x00-\x1f]', "_", name)
    name = re.sub(r"\s+", " ", name).strip(" .")
    return name[:120] or "документ"


@dataclass
class FillPlan:
    column_of: dict[str, str]                    # имя в шаблоне -> колонка таблицы
    missing: list[str] = field(default_factory=list)
    bad_formats: list[str] = field(default_factory=list)


def plan_fields(placeholders: list[tuple[str, str | None]], columns: list[str]) -> FillPlan:
    by_key = {norm_key(c): c for c in columns}
    plan = FillPlan({})
    for name, fmt in placeholders:
        col = by_key.get(norm_key(name))
        if col is None:
            if name not in plan.missing:
                plan.missing.append(name)
        else:
            plan.column_of[name] = col
        if fmt and fmt.lower() not in FORMATS and fmt not in plan.bad_formats:
            plan.bad_formats.append(fmt)
    return plan


def fill_documents(template: Path, data: Path, out_dir: Path, *, name_pattern: str | None = None,
                   sheet: str | None = None, allow_missing: bool = False, to_pdf: bool = False) -> tuple[list[Path], Report]:
    template, data, out_dir = Path(template), Path(data), Path(out_dir)
    if not template.exists():
        raise DocToolsError(f"Шаблон не найден: {template}")
    if template.suffix.lower() != ".docx":
        raise DocToolsError(f"{template.name}: шаблон должен быть .docx (старый .doc — сначала convert --to docx)")
    table = read_table(data, sheet)
    try:
        probe = Document(str(template))
    except Exception as exc:
        raise DocToolsError(f"{template.name}: не удалось открыть как документ Word ({exc})") from exc
    placeholders = find_placeholders(probe)
    if not placeholders:
        raise DocToolsError(f"{template.name}: в шаблоне нет плейсхолдеров вида {{{{Поле}}}}")
    name_pattern = name_pattern or template.stem + "_{{№}}"
    name_fields = [(m.group(1), m.group(2)) for m in PLACEHOLDER.finditer(name_pattern)]
    plan = plan_fields(placeholders + [f for f in name_fields if f[0] != "№"], table.columns)
    if plan.bad_formats:
        raise DocToolsError(f"Неизвестный формат: {', '.join(plan.bad_formats)}. Есть: {', '.join(FORMATS)}")
    if plan.missing and not allow_missing:
        raise DocToolsError(f"В таблице нет колонок для плейсхолдеров: {', '.join(plan.missing)}. "
                            f"Колонки таблицы: {', '.join(table.columns)}. "
                            "Исправьте названия или запустите с --allow-missing (поля останутся пустыми)")

    rep = Report(f"Заполнение шаблона: {template.name} × {data.name}")
    out_dir.mkdir(parents=True, exist_ok=True)
    made: list[Path] = []
    empty_fields: list[list[Any]] = []
    skipped = 0
    used_names: set[str] = set()
    for n, (row, rn) in enumerate(zip(table.rows, table.row_numbers), start=1):
        if all(is_empty(v) for v in row):
            skipped += 1
            continue
        values = dict(zip(table.columns, row))
        problems: list[str] = []

        def render(name: str, fmt: str | None, _values: dict = values, _n: int = n, _problems: list = problems) -> str:
            if name == "№":
                return str(_n)
            col = plan.column_of.get(name)
            v = _values.get(col) if col else None
            if is_empty(v):
                if name not in _problems:
                    _problems.append(name)
                return ""
            if fmt:
                try:
                    return FORMATS[fmt.lower()](v)
                except (ValueError, TypeError):
                    _problems.append(f"{name}|{fmt}: «{v}» не подходит под формат")
                    return default_text(v)
            return default_text(v)

        doc = Document(str(template))
        for p in _paragraphs(doc):
            _replace_in_paragraph(p, render)
        fname = _safe_filename(PLACEHOLDER.sub(lambda m: render(m.group(1), m.group(2)), name_pattern))
        base, k = fname, 2
        while fname.lower() in used_names:
            fname = f"{base}_{k}"
            k += 1
        used_names.add(fname.lower())
        path = out_dir / f"{fname}.docx"
        try:
            doc.save(str(path))
        except PermissionError as exc:
            raise DocToolsError(f"Не удалось записать {path.name}: файл открыт в Word?") from exc
        made.append(path)
        if problems:
            empty_fields.append([rn, path.name, ", ".join(problems)])

    rep.stat("Строк в таблице", len(table.rows))
    rep.stat("Документов создано", len(made))
    rep.stat("Пустых строк пропущено", skipped)
    rep.stat("Плейсхолдеров в шаблоне", len(placeholders))
    rep.info("Плейсхолдеры: " + ", ".join("{{" + a + (f"|{b}" if b else "") + "}}" for a, b in placeholders))
    unused = [c for c in table.columns if c not in plan.column_of.values()]
    if unused:
        rep.info(f"Колонки таблицы, не использованные в шаблоне: {', '.join(unused)}")
    if plan.missing:
        rep.warn(f"Нет в таблице (оставлены пустыми): {', '.join(plan.missing)}")
    if empty_fields:
        rep.warn(f"Документов с пустыми полями: {len(empty_fields)} — проверьте перед отправкой")
        rep.table("Пустые или неподходящие значения", ["Строка в таблице", "Документ", "Поля"], empty_fields)
    for p in made:
        rep.output(p)

    if to_pdf and made:
        from .word_com import WordSession

        done = 0
        with WordSession() as word:
            for p in made:
                try:
                    rep.output(word.to_pdf(p, p.with_suffix(".pdf")))
                    done += 1
                except DocToolsError as exc:
                    rep.error(str(exc))
        rep.info(f"PDF сделаны через Microsoft Word: {done} из {len(made)}")
        if word.killed:
            rep.warn(f"Word не закрылся сам — завершён принудительно (PID {', '.join(map(str, word.killed))})")
    return made, rep
