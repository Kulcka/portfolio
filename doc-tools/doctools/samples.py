"""Генератор демо-набора: документы вымышленной компании ООО «СеверСнаб-Демо».

Все данные вымышленные (152-ФЗ): имена составлены из распространённых имён и
фамилий, телефоны вида +7 (900) 000-xx-xx, почта — в зарезервированных доменах
example.com/.org/.net, ИНН начинаются с несуществующего кода региона «00».

Используется и для examples/, и как фикстуры тестов: ожидаемые значения
(число строк, суммы) экспортируются константами.
"""

from __future__ import annotations

import csv
import io
import random
from dataclasses import dataclass
from datetime import date, datetime
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt, RGBColor
from fpdf import FPDF, FontFace
from openpyxl import Workbook
from openpyxl.styles import Font

from .normalize import _W10, _ctrl, format_money

FONT_DIR = Path(r"C:\Windows\Fonts")
COMPANY = "ООО «СеверСнаб-Демо»"
COMPANY_LINE = f"{COMPANY} · г. Демоград, ул. Примерная, 1 · +7 (900) 000-00-00 · info@example.com"


def demo_inn(seed: int) -> str:
    """Вымышленный ИНН организации с правильной контрольной цифрой и кодом региона 00."""
    body = f"00{seed:07d}"
    return body + str(_ctrl(body, _W10))


# ============================================================ прайс (3 стр.)
PRICE_SECTIONS: dict[str, list[tuple[str, str, str]]] = {
    "Крепёж": [
        ("Болт М6×20 оцинкованный, DIN 933", "уп.", "по 100 шт."),
        ("Болт М8×30 оцинкованный, DIN 933", "уп.", ""),
        ("Болт М10×40 оцинкованный, DIN 933", "уп.", ""),
        ("Гайка М6 оцинкованная, DIN 934", "уп.", ""),
        ("Гайка М8 оцинкованная, DIN 934", "уп.", ""),
        ("Гайка М10 самоконтрящаяся с нейлоновым кольцом, DIN 985", "уп.", ""),
        ("Шайба плоская М6, DIN 125", "уп.", ""),
        ("Шайба плоская М8, DIN 125", "уп.", ""),
        ("Шайба гроверная М8, DIN 127", "уп.", ""),
        ("Саморез по дереву 3,5×35, жёлтый цинк", "уп.", ""),
        ("Саморез по металлу 4,2×16 с прессшайбой, острый", "уп.", ""),
        ("Саморез кровельный 4,8×29 с EPDM-шайбой, RAL 3005", "уп.", ""),
        ("Анкер клиновой 10×100", "шт.", ""),
        ("Анкерный болт с гайкой 12×129", "шт.", ""),
        ("Дюбель-гвоздь 6×40 с потайным бортиком", "уп.", ""),
        ("Шпилька резьбовая М10×1000, DIN 975", "шт.", ""),
        ("Хомут червячный 16–25 мм, нержавеющая сталь", "шт.", ""),
    ],
    "Инструмент": [
        ("Дрель-шуруповёрт аккумуляторная 18 В, 2 АКБ, кейс", "шт.", ""),
        ("Перфоратор SDS-plus 800 Вт", "шт.", ""),
        ("Угловая шлифмашина 125 мм, 900 Вт", "шт.", ""),
        ("Набор бит 32 предмета в кейсе", "набор", ""),
        ("Набор свёрл по металлу HSS 1–10 мм, 19 шт.", "набор", ""),
        ("Бур SDS-plus 8×160", "шт.", ""),
        ("Бур SDS-plus 10×210", "шт.", ""),
        ("Круг отрезной по металлу 125×1,2", "шт.", ""),
        ("Рулетка 5 м с магнитным зацепом", "шт.", ""),
        ("Уровень строительный 600 мм, 3 глазка", "шт.", ""),
        ("Молоток слесарный 500 г, фибергласовая рукоять", "шт.", ""),
        ("Ключ разводной 250 мм", "шт.", ""),
        ("Набор головок 1/2\" 24 предмета", "набор", ""),
        ("Лазерный нивелир с треногой, зелёный луч", "шт.", ""),
        ("Строительный фен 2000 Вт с насадками", "шт.", ""),
        ("Тиски слесарные поворотные 125 мм", "шт.", ""),
    ],
    "Электрика": [
        ("Кабель ВВГнг-LS 3×2,5 (бухта 100 м)", "бухта", ""),
        ("Кабель ВВГнг-LS 3×1,5 (бухта 100 м)", "бухта", ""),
        ("Автоматический выключатель 1P 16А, характеристика C", "шт.", ""),
        ("Автоматический выключатель 1P 25А, характеристика C", "шт.", ""),
        ("УЗО 2P 40А 30мА", "шт.", ""),
        ("Розетка двойная с заземлением, белая", "шт.", ""),
        ("Выключатель одноклавишный, белый", "шт.", ""),
        ("Коробка распаячная 100×100×50, IP55", "шт.", ""),
        ("Гофротруба ПВХ 20 мм с зондом (бухта 50 м)", "бухта", ""),
        ("Светильник светодиодный 36 Вт, 600×600, 4000К", "шт.", ""),
        ("Лампа светодиодная E27 12 Вт, 4000К", "шт.", ""),
        ("Удлинитель 5 розеток, 5 м, с выключателем", "шт.", ""),
        ("Щит распределительный на 12 модулей, встраиваемый", "шт.", ""),
        ("Клеммы соединительные 3×0,2–4 мм², 50 шт.", "уп.", ""),
        ("Изолента ПВХ 19 мм × 20 м, синяя", "шт.", ""),
    ],
}
PRICE_ON_REQUEST = {"Лазерный нивелир с треногой, зелёный луч", "Щит распределительный на 12 модулей, встраиваемый"}
PRICE_UNDER_ORDER = {"Анкерный болт с гайкой 12×129", "Строительный фен 2000 Вт с насадками"}
PRICE_ITEMS = sum(len(v) for v in PRICE_SECTIONS.values())          # строк товаров
DISCOUNTS = [("30 000,00", "3"), ("100 000,00", "5"), ("300 000,00", "7"), ("1 000 000,00", "10")]


class _DemoPDF(FPDF):
    footer_text = ""

    def multi_cell(self, *args, **kwargs):  # type: ignore[override]
        kwargs.setdefault("new_x", "LMARGIN")
        kwargs.setdefault("new_y", "NEXT")
        return super().multi_cell(*args, **kwargs)

    def header(self) -> None:
        self.set_font("Arial", size=8)
        self.set_text_color(110, 110, 110)
        self.cell(0, 5, COMPANY_LINE, align="R")
        self.ln(9)
        self.set_text_color(0, 0, 0)

    def footer(self) -> None:
        self.set_y(-12)
        self.set_font("Arial", size=8)
        self.set_text_color(110, 110, 110)
        self.cell(0, 5, f"{self.footer_text} · Стр. {self.page_no()} из {{nb}}", align="C")


def _pdf(footer: str) -> _DemoPDF:
    if not (FONT_DIR / "arial.ttf").exists():
        raise RuntimeError("Для генерации демо-PDF нужен шрифт Arial (C:\\Windows\\Fonts\\arial.ttf)")
    pdf = _DemoPDF(format="A4")
    pdf.footer_text = footer
    pdf.add_font("Arial", "", str(FONT_DIR / "arial.ttf"))
    pdf.add_font("Arial", "B", str(FONT_DIR / "arialbd.ttf"))
    pdf.add_font("Arial", "I", str(FONT_DIR / "ariali.ttf"))
    pdf.set_auto_page_break(True, margin=18)
    pdf.set_margins(14, 12, 14)
    return pdf


def _price_rows(rnd: random.Random) -> list[tuple[str, list[str]]]:
    rows = []
    n = 0
    for section, items in PRICE_SECTIONS.items():
        rows.append((section, []))
        for name, unit, _ in items:
            n += 1
            code = f"СС-{n:04d}"
            if name in PRICE_ON_REQUEST:
                price = "по запросу"
            else:
                price = format_money(round(rnd.uniform(40, 25000), -1) + rnd.choice([0, 0.5, 0.9]))
            stock = "под заказ" if name in PRICE_UNDER_ORDER else str(rnd.randint(0, 900))
            arrival = date(2026, rnd.randint(7, 9), rnd.randint(1, 28)).strftime("%d.%m.%Y")
            rows.append(("", [code, name, unit, price, stock, arrival]))
    return rows


def make_price_pdf(path: Path) -> Path:
    rnd = random.Random(42)
    pdf = _pdf("Прайс-лист на 01.09.2026")
    pdf.add_page()
    pdf.set_font("Arial", "B", 18)
    pdf.cell(0, 10, "Прайс-лист", new_x="LMARGIN", new_y="NEXT")
    pdf.set_font("Arial", "", 11)
    pdf.cell(0, 6, "на 01.09.2026, цены в рублях с НДС 20 %", new_x="LMARGIN", new_y="NEXT")
    pdf.ln(3)
    pdf.set_font("Arial", "", 10)
    pdf.multi_cell(0, 5, "Уважаемые партнёры! Направляем актуальные цены на крепёж, инструмент и "
                         "электротовары. Цены действуют при заказе со склада в Демограде. Позиции с "
                         "пометкой «под заказ» поставляются в течение 5–7 рабочих дней, цены «по запросу» "
                         "уточняйте у менеджера.")
    pdf.ln(4)
    head = FontFace(emphasis="BOLD", fill_color=(221, 235, 247))
    sect = FontFace(emphasis="BOLD", fill_color=(242, 242, 242))
    pdf.set_font("Arial", "", 9)
    with pdf.table(col_widths=(20, 78, 13, 26, 20, 25), line_height=5, repeat_headings=1,
                   headings_style=head, text_align=("LEFT", "LEFT", "CENTER", "RIGHT", "RIGHT", "CENTER")) as table:
        table.row(["Артикул", "Наименование", "Ед.", "Цена, ₽", "Остаток", "Поступление"])
        for section, cells in _price_rows(rnd):
            row = table.row()
            if section:
                row.cell(section, colspan=6, style=sect)
            else:
                for c in cells:
                    row.cell(c)
    pdf.add_page()
    pdf.set_font("Arial", "B", 14)
    pdf.cell(0, 8, "Скидки при заказе от суммы", new_x="LMARGIN", new_y="NEXT")
    pdf.ln(2)
    pdf.set_font("Arial", "", 10)
    with pdf.table(col_widths=(60, 40), width=100, align="LEFT", line_height=6, headings_style=head,
                   text_align=("RIGHT", "RIGHT")) as table:
        table.row(["Сумма заказа от, ₽", "Скидка, %"])
        for a, b in DISCOUNTS:
            table.row([a, b])
    pdf.ln(6)
    pdf.set_font("Arial", "B", 14)
    pdf.cell(0, 8, "Условия поставки", new_x="LMARGIN", new_y="NEXT")
    pdf.set_font("Arial", "", 10)
    pdf.multi_cell(0, 5, "Отгрузка со склада производится в день оплаты при поступлении денег до 14:00. "
                         "Счёт действителен три банковских дня. Скидки не суммируются с акционными ценами "
                         "и применяются ко всей сумме заказа.")
    pdf.ln(3)
    for text in ("Доставка по Демограду — бесплатно при заказе от 15 000 ₽.",
                 "Доставка в другие города — транспортной компанией по тарифам перевозчика.",
                 "Самовывоз — пн–пт с 9:00 до 18:00, сб с 10:00 до 15:00."):
        pdf.multi_cell(0, 5, f"•  {text}")
    pdf.ln(3)
    pdf.set_font("Arial", "B", 11)
    pdf.cell(0, 7, "Возврат и обмен", new_x="LMARGIN", new_y="NEXT")
    pdf.set_font("Arial", "", 10)
    pdf.multi_cell(0, 5, "Возврат товара надлежащего качества возможен в течение 14 дней при сохранении "
                         "упаковки и товарного вида. Инструмент с признаками эксплуатации возврату не "
                         "подлежит, гарантийные случаи рассматриваются сервисным центром производителя.")
    path.parent.mkdir(parents=True, exist_ok=True)
    pdf.output(str(path))
    return path


# ============================================================ выписка (2 стр.)
COUNTERPARTIES = ["ООО «Ромашка-Демо»", "ИП Кузнецов Д. П.", "ООО «ТоргПример»", "АО «Демо-Логистик»",
                  "ООО «Вектор-Тест»", "ИП Соколова М. А.", "ООО «СтройОбразец»"]
PURPOSES_IN = ["Оплата по счёту № {n} от {d}, в т.ч. НДС 20 %", "Предоплата по договору № {n}",
               "Оплата за крепёж по счёту № {n}"]
PURPOSES_OUT = ["Оплата поставщику по договору № {n}, в т.ч. НДС 20 %", "Аренда склада за сентябрь 2026",
                "Транспортные услуги по акту № {n}", "Комиссия банка за ведение счёта"]
STATEMENT_OPS = 40
STATEMENT_OPENING = 250000.00


def statement_operations() -> list[tuple[date, str, str, str, float, float]]:
    rnd = random.Random(7)
    ops = []
    for i in range(STATEMENT_OPS):
        d = date(2026, 9, 1 + i * 15 // STATEMENT_OPS)
        income = rnd.random() < 0.55
        amount = round(rnd.uniform(1500, 90000), 2)
        n = rnd.randint(100, 999)
        text = (rnd.choice(PURPOSES_IN) if income else rnd.choice(PURPOSES_OUT)).format(n=n, d=d.strftime("%d.%m.%Y"))
        cp = "АО «Банк Демо»" if "Комиссия" in text else rnd.choice(COUNTERPARTIES)
        if "Комиссия" in text:
            amount = round(amount / 100, 2)  # комиссия банка — сотни рублей, не десятки тысяч
        ops.append((d, str(1000 + i), cp, text, amount if income else 0.0, 0.0 if income else amount))
    return ops


def make_statement_pdf(path: Path) -> Path:
    ops = statement_operations()
    pdf = _pdf("Выписка по счёту 40702810000000000001")
    pdf.add_page()
    pdf.set_font("Arial", "B", 14)
    pdf.cell(0, 8, "Выписка по расчётному счёту", new_x="LMARGIN", new_y="NEXT")
    pdf.set_font("Arial", "", 10)
    for line in (f"Клиент: {COMPANY}, ИНН {demo_inn(1234)}",
                 "Счёт: 40702810000000000001 в АО «Банк Демо» (демонстрационные данные)",
                 "Период: 01.09.2026 — 15.09.2026",
                 f"Входящий остаток: {format_money(STATEMENT_OPENING)} ₽"):
        pdf.cell(0, 5.5, line, new_x="LMARGIN", new_y="NEXT")
    pdf.ln(3)
    pdf.set_font("Arial", "", 8.5)
    head = FontFace(emphasis="BOLD")
    total_in = sum(o[4] for o in ops)
    total_out = sum(o[5] for o in ops)
    with pdf.table(col_widths=(18, 13, 36, 67, 24, 24), line_height=4.5, repeat_headings=0,
                   borders_layout="HORIZONTAL_LINES", headings_style=head,
                   text_align=("LEFT", "LEFT", "LEFT", "LEFT", "RIGHT", "RIGHT")) as table:
        table.row(["Дата", "№ док.", "Контрагент", "Назначение платежа", "Поступление", "Списание"])
        for d, n, cp, text, inc, out in ops:
            table.row([d.strftime("%d.%m.%Y"), n, cp, text,
                       format_money(inc) if inc else "", format_money(out) if out else ""])
        table.row(["", "", "", "Обороты за период", format_money(total_in), format_money(total_out)])
    pdf.ln(4)
    pdf.set_font("Arial", "", 10)
    closing = STATEMENT_OPENING + total_in - total_out
    pdf.cell(0, 6, f"Исходящий остаток: {format_money(closing)} ₽", new_x="LMARGIN", new_y="NEXT")
    path.parent.mkdir(parents=True, exist_ok=True)
    pdf.output(str(path))
    return path


# ============================================================ скан без текста
def make_scan_pdf(path: Path) -> Path:
    """PDF-«скан»: страница — картинка, текстового слоя нет."""
    from PIL import Image, ImageDraw, ImageFilter, ImageFont

    img = Image.new("L", (1240, 1754), 250)
    draw = ImageDraw.Draw(img)
    big = ImageFont.truetype(str(FONT_DIR / "arialbd.ttf"), 44)
    small = ImageFont.truetype(str(FONT_DIR / "arial.ttf"), 30)
    draw.text((140, 160), "ДОГОВОР ПОСТАВКИ № 17-Д", font=big, fill=20)
    draw.text((140, 240), "г. Демоград                                        01 сентября 2026 г.", font=small, fill=30)
    y = 330
    for line in ["ООО «СеверСнаб-Демо», именуемое в дальнейшем «Поставщик»,",
                 "и ООО «Ромашка-Демо», именуемое в дальнейшем «Покупатель»,",
                 "заключили настоящий договор о нижеследующем.",
                 "", "1. Предмет договора",
                 "1.1. Поставщик обязуется передать, а Покупатель — принять и оплатить",
                 "товар по спецификации (Приложение № 1)."]:
        draw.text((140, y), line, font=small, fill=35)
        y += 48
    img = img.rotate(0.7, fillcolor=250).filter(ImageFilter.GaussianBlur(0.6))
    buf = io.BytesIO()
    img.save(buf, format="PNG")
    buf.seek(0)
    pdf = FPDF(format="A4")
    pdf.add_page()
    pdf.image(buf, x=0, y=0, w=210, h=297)
    path.parent.mkdir(parents=True, exist_ok=True)
    pdf.output(str(path))
    return path


# ============================================================ таблицы клиентов
def _save_wb(wb: Workbook, path: Path) -> Path:
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    return path


INN = {k: demo_inn(v) for k, v in {"ivanova": 1001, "petrov": 2002, "smirnova": 3003, "sokolova": 5005,
                                    "morozova": 6006, "volkov": 7007, "fedorov": 8008}.items()}
BAD_INN = "0012345678"  # контрольная цифра не сходится (специально для демо)


def make_sales_xlsx(path: Path) -> Path:
    """Отдел продаж: свои названия колонок, «грязные» значения, дубли, пустая строка."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Клиенты"
    ws.append(["ФИО", "Телефон", "E-mail", "Город", "ИНН", "Дата обращения", "Сумма заказа"])
    rows = [
        ["Иванова Анна Сергеевна", "8 (900) 000-12-34", "anna.ivanova@example.com", "Демоград", int(INN["ivanova"]), datetime(2026, 9, 1), 12500.5],
        ["  петров  иван ", "+7 900 000 23 45", "PETROV@EXAMPLE.COM", "демоград", INN["petrov"], "05.09.2026", "8 300,00"],
        ["Смирнова Ольга Викторовна", "89000003456", "o.smirnova@example.org", "Нижний Пример", INN["smirnova"], "2026-09-03", "15 000 руб."],
        ["Кузнецов Дмитрий Павлович", "900-000-45-67", "kuznetsov@example.net", "Примерск", BAD_INN, datetime(2026, 9, 4), 4200],
        ["Соколова Мария", "+7 (900) 000-56-78", "maria.s@example.com", "Демоград", INN["sokolova"], "7 сентября 2026", "1,250.75"],
        [None, None, None, None, None, None, None],
        ["Попов Алексей Игоревич", "000-67-89", "popov@example.com", "Примерск", None, "08.09.2026", 3100],
        ["Васильева Елена", "8 900 000 78 01", "vasilieva@example", "Демоград", None, "09.09.2026", "2\u00a0400"],
        ["Смирнова Ольга Викторовна", "89000003456", "o.smirnova@example.org", "Нижний Пример", INN["smirnova"], "2026-09-03", "15 000 руб."],
        ["Новиков Сергей", "+7 900 000 89 12", "novikov@example.com", "Нижний Пример", None, "10.09.2026", 7700],
    ]
    for r in rows:
        ws.append(r)
    ws["A1"].font = Font(bold=True)
    return _save_wb(wb, path)


def make_crm_xlsx(path: Path) -> Path:
    """Выгрузка CRM: заголовок отчёта над таблицей, синонимы колонок, лишняя колонка."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Выгрузка"
    ws.append(["Выгрузка из CRM на 15.09.2026"])
    ws.append([])
    ws.append(["Контактное лицо", "Тел.", "Почта", "Населённый пункт", "ИНН организации", "Дата", "Сумма, ₽", "Комментарий"])
    rows = [
        ["Смирнова Ольга", "+79000003456", "o.smirnova@example.org", "Нижний Пример", INN["smirnova"], "11.09.2026", 5000, "повторное обращение"],
        ["Кузнецов Дмитрй Павлович", "8 900 000 45 67", "kuznetsov@example.net", "Примерск", BAD_INN, "12.09.2026", "6 100,00", ""],
        ["Морозова Татьяна", "8(900)000-90-23", "morozova@example.org", "Демоград", INN["morozova"], "12.09.2026", "9 999,99", "нужен счёт"],
        ["Волков Андрей", "+7 900 000-01-34", "volkov@example.net", "примерск", INN["volkov"], "13.09.2026", 18000, ""],
        ["Иванова Анна Сергеевна", "+7 (900) 000-12-34", None, "Демоград", INN["ivanova"], "01.09.2026", 12500.5, "просила счёт на e-mail"],
        ["Лебедева Ирина", "+7 900 000 12 45", "lebedeva@example.com", "Демоград", None, "14.09.2026", "2 750", ""],
    ]
    for r in rows:
        ws.append(r)
    ws["A1"].font = Font(bold=True, size=12)
    return _save_wb(wb, path)


def make_site_csv(path: Path) -> Path:
    """Заявки с сайта: CSV в Windows-1251 с «;», даты ISO, суммы с точкой."""
    rows = [
        ["Имя", "Номер телефона", "Email", "Город", "Дата заявки", "Сумма"],
        ["Федоров Павел", "8 900 000 78 90", "FEDOROV@EXAMPLE.COM", "Демоград", "2026-09-10", "3200.00"],
        ["петров иван", "+7(900)0002345", "petrov@example.com", "Демоград", "12.09.26", "8300"],
        ["Орлова Светлана", "89000008901", "orlova@example.org", "Нижний Пример", "2026-09-13", "4 450.50"],
        ["Зайцев Никита", "8-900-000-99-01", "zaitsev@example.net", "Примерск", "2026-09-14", "12000"],
    ]
    buf = io.StringIO()
    csv.writer(buf, delimiter=";", lineterminator="\r\n").writerows(rows)
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_bytes(buf.getvalue().encode("cp1251"))
    return path


# ============================================================ шаблон письма
RECIPIENTS = [
    ["Иванова Анна Сергеевна", "Анна Сергеевна", "Уважаемая", "генеральному директору", "ООО «Ромашка-Демо»", "anna.ivanova@example.com", "17-Д", date(2025, 10, 1), 1250000],
    ["Петров Иван Николаевич", "Иван Николаевич", "Уважаемый", "директору по закупкам", "ООО «ТоргПример»", "petrov@example.com", "23-Д", date(2025, 11, 15), 384500.5],
    ["Смирнова Ольга Викторовна", "Ольга Викторовна", "Уважаемая", "руководителю отдела снабжения", "АО «Демо-Логистик»", "o.smirnova@example.org", "31-Д", date(2026, 1, 20), 96000],
    ["Кузнецов Дмитрий Павлович", "Дмитрий Павлович", "Уважаемый", None, "ИП Кузнецов Д. П.", "kuznetsov@example.net", "44-Д", date(2026, 3, 2), 45210.75],
    ["Морозова Татьяна Алексеевна", "Татьяна Алексеевна", "Уважаемая", "главному инженеру", "ООО «СтройОбразец»", "morozova@example.org", "52-Д", date(2026, 4, 11), 2001001],
]
RECIPIENT_COLUMNS = ["ФИО", "Имя и отчество", "Обращение", "Должность", "Компания", "Email", "Номер договора", "Дата договора", "Сумма"]


def make_recipients_xlsx(path: Path) -> Path:
    wb = Workbook()
    ws = wb.active
    ws.title = "Адресаты"
    ws.append(RECIPIENT_COLUMNS)
    for r in RECIPIENTS:
        ws.append(r)
    for c in ws[1]:
        c.font = Font(bold=True)
    for cell in ws["H"][1:]:
        cell.number_format = "DD.MM.YYYY"
    return _save_wb(wb, path)


def _split_runs(paragraph, pieces: list[tuple[str, bool]]) -> None:
    for text, bold in pieces:
        run = paragraph.add_run(text)
        run.bold = bold


def make_letter_template(path: Path) -> Path:
    """Шаблон письма с плейсхолдерами {{поле}}; часть плейсхолдеров намеренно разбита на
    несколько фрагментов текста — так Word сохраняет их после правок, и заполнение должно это выдерживать."""
    doc = Document()
    st = doc.styles["Normal"]
    st.font.name = "Arial"
    st.font.size = Pt(11)
    sec = doc.sections[0]
    hp = sec.header.paragraphs[0]
    hp.text = COMPANY_LINE
    hp.runs[0].font.size = Pt(8)
    hp.runs[0].font.color.rgb = RGBColor(0x70, 0x70, 0x70)
    fp = sec.footer.paragraphs[0]
    fp.text = "Письмо подготовлено для {{Компания}} · исх. № {{Номер договора}}/П"
    fp.runs[0].font.size = Pt(8)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    p.add_run("{{Должность}}\n{{Компания}}\n{{ФИО|инициалы}}\n{{Email}}")
    p = doc.add_paragraph()
    p.add_run("Исх. № {{Номер договора}}/П от 27.09.2026")
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _split_runs(p, [("О продлении договора № ", True), ("{{Номер ", True), ("договора}}", True)])
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _split_runs(p, [("{{", False), ("Обращение", False), ("}} {{Имя и отчество}}!", False)])
    doc.add_paragraph(
        "Благодарим вас за сотрудничество. Срок действия договора поставки № {{Номер договора}} "
        "от {{Дата договора|дата}} истекает 31.12.2026. Предлагаем продлить его на 2027 год на прежних "
        "условиях. Объём поставок по договору за период составил {{Сумма|деньги}} ₽ "
        "({{Сумма|прописью}}).")
    doc.add_paragraph("Основные условия продления:")
    table = doc.add_table(rows=4, cols=2)
    table.style = "Table Grid"
    for i, (a, b) in enumerate([("Покупатель", "{{Компания}}"), ("Договор", "№ {{Номер договора}} от {{Дата договора|дата}}"),
                                ("Объём за период, ₽", "{{Сумма|деньги}}"), ("Новый срок действия", "до 31.12.2027")]):
        table.cell(i, 0).text = a
        table.cell(i, 1).text = b
    doc.add_paragraph()
    doc.add_paragraph("Просим до 15.10.2026 сообщить о своём решении ответным письмом.")
    doc.add_paragraph("С уважением,\nкоммерческий директор\n" + COMPANY + "                     П. П. Примеров")
    path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(path)
    return path


# ============================================================ всё сразу
@dataclass
class SampleSet:
    price_pdf: Path
    statement_pdf: Path
    scan_pdf: Path
    sales_xlsx: Path
    crm_xlsx: Path
    site_csv: Path
    recipients_xlsx: Path
    letter_docx: Path


def generate(folder: Path) -> SampleSet:
    folder = Path(folder)
    folder.mkdir(parents=True, exist_ok=True)
    return SampleSet(
        price_pdf=make_price_pdf(folder / "прайс_СеверСнаб.pdf"),
        statement_pdf=make_statement_pdf(folder / "выписка_сентябрь.pdf"),
        scan_pdf=make_scan_pdf(folder / "скан_договора.pdf"),
        sales_xlsx=make_sales_xlsx(folder / "клиенты_отдел_продаж.xlsx"),
        crm_xlsx=make_crm_xlsx(folder / "выгрузка_CRM.xlsx"),
        site_csv=make_site_csv(folder / "заявки_с_сайта.csv"),
        recipients_xlsx=make_recipients_xlsx(folder / "адресаты.xlsx"),
        letter_docx=make_letter_template(folder / "шаблон_письма.docx"),
    )
