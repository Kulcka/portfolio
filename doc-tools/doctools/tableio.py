"""Чтение и запись таблиц: XLSX, CSV, JSON. Оформленный Excel на выходе.

Внутреннее представление — Table: список названий колонок и список строк
(значения Python: str, int, float, date, datetime, None). pandas не используем
намеренно: он молча превращает ИНН и телефоны в числа и теряет ведущие нули.
"""

from __future__ import annotations

import csv
import io
import json
import re
from dataclasses import dataclass, field
from datetime import date, datetime
from pathlib import Path
from typing import Any, Iterable

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from . import DocToolsError
from .normalize import clean_spaces, is_empty, parse_date, parse_number

TABLE_EXT = {".xlsx", ".xlsm", ".csv", ".json"}


@dataclass
class Table:
    columns: list[str]
    rows: list[list[Any]]
    name: str = ""                                   # «файл.xlsx / Лист1»
    row_numbers: list[int] = field(default_factory=list)  # номер строки в исходнике

    def __post_init__(self) -> None:
        if not self.row_numbers:
            self.row_numbers = list(range(2, len(self.rows) + 2))

    def col(self, name: str) -> int:
        return self.columns.index(name)


# ================================================================== чтение
def read_tables(path: Path, sheet: str | None = None, header_row: int | None = None) -> list[Table]:
    """Все таблицы файла (для Excel — по листу на таблицу, пустые листы пропускаются)."""
    path = Path(path)
    if not path.exists():
        raise DocToolsError(f"Файл не найден: {path}")
    ext = path.suffix.lower()
    if ext in (".xlsx", ".xlsm"):
        return _read_xlsx(path, sheet, header_row)
    if ext == ".csv":
        return [_grid_to_table(_read_csv_grid(path), path.name, header_row)]
    if ext == ".json":
        return _read_json(path)
    if ext == ".xls":
        raise DocToolsError(f"{path.name}: старый формат .xls не читается напрямую — "
                            "пересохраните в .xlsx (или конвертируем через Excel по согласованию).")
    raise DocToolsError(f"{path.name}: неподдерживаемый формат таблицы {ext} (нужно .xlsx, .csv или .json)")


def read_table(path: Path, sheet: str | None = None, header_row: int | None = None) -> Table:
    """Одна таблица: указанный лист или первый непустой."""
    tables = read_tables(path, sheet, header_row)
    if not tables:
        raise DocToolsError(f"{Path(path).name}: в файле нет данных")
    return tables[0]


def _read_xlsx(path: Path, sheet: str | None, header_row: int | None) -> list[Table]:
    try:
        wb = load_workbook(path, data_only=True, read_only=True)
    except Exception as exc:  # повреждённый файл, защищён паролем и т.п.
        raise DocToolsError(f"{path.name}: не удалось открыть как Excel ({exc})") from exc
    try:
        names = wb.sheetnames
        if sheet is not None:
            if sheet not in names:
                raise DocToolsError(f"{path.name}: нет листа «{sheet}». Есть: {', '.join(names)}")
            names = [sheet]
        tables = []
        for name in names:
            grid = [list(r) for r in wb[name].iter_rows(values_only=True)]
            if any(not is_empty(v) for row in grid for v in row):
                label = path.name if len(wb.sheetnames) == 1 else f"{path.name} / {name}"
                tables.append(_grid_to_table(grid, label, header_row))
        return tables
    finally:
        wb.close()


def _read_csv_grid(path: Path) -> list[list[Any]]:
    raw = path.read_bytes()
    for enc in ("utf-8-sig", "cp1251"):
        try:
            text = raw.decode(enc)
            break
        except UnicodeDecodeError:
            continue
    else:
        raise DocToolsError(f"{path.name}: не удалось определить кодировку (ожидалась UTF-8 или Windows-1251)")
    sample = text[:20000]
    try:
        dialect = csv.Sniffer().sniff(sample, delimiters=";,\t|")
        delim = dialect.delimiter
    except csv.Error:
        delim = ";" if sample.count(";") >= sample.count(",") else ","
    return [row for row in csv.reader(io.StringIO(text), delimiter=delim)]


def _read_json(path: Path) -> list[Table]:
    try:
        data = json.loads(path.read_text(encoding="utf-8-sig"))
    except json.JSONDecodeError as exc:
        raise DocToolsError(f"{path.name}: некорректный JSON ({exc})") from exc
    if isinstance(data, list):
        return [_records_to_table(data, path.name)]
    if isinstance(data, dict) and all(isinstance(v, list) for v in data.values()):
        return [_records_to_table(v, f"{path.name} / {k}") for k, v in data.items()]
    raise DocToolsError(f"{path.name}: ожидался список записей [{{...}}, ...] или {{\"лист\": [...]}}")


def _records_to_table(records: list[Any], name: str) -> Table:
    columns: list[str] = []
    for rec in records:
        if not isinstance(rec, dict):
            raise DocToolsError(f"{name}: элементы списка должны быть объектами {{\"колонка\": значение}}")
        for k in rec:
            if k not in columns:
                columns.append(k)
    rows = []
    for rec in records:
        row = []
        for c in columns:
            v = rec.get(c)
            row.append(json.dumps(v, ensure_ascii=False) if isinstance(v, (dict, list)) else v)
        rows.append(row)
    return Table(columns, rows, name)


def find_header_row(grid: list[list[Any]], scan: int = 15) -> int:
    """Индекс строки заголовка: первая «широкая» строка из одних текстов.

    Пропускает заголовки отчётов над таблицей («Выгрузка из CRM на …») и пустые строки.
    """
    head = grid[:scan]
    widths = [sum(not is_empty(v) for v in row) for row in head]
    if not widths or max(widths) == 0:
        return 0
    need = max(2, int(max(widths) * 0.6))
    for i, row in enumerate(head):
        vals = [v for v in row if not is_empty(v)]
        if len(vals) >= need and all(isinstance(v, str) and parse_number(v) is None for v in vals):
            return i
    return next(i for i, w in enumerate(widths) if w)


def _grid_to_table(grid: list[list[Any]], name: str, header_row: int | None) -> Table:
    if not grid:
        return Table([], [], name)
    h = (header_row - 1) if header_row else find_header_row(grid)
    width = max(len(r) for r in grid)
    grid = [list(r) + [None] * (width - len(r)) for r in grid]
    header = grid[h]
    # обрезаем пустые колонки справа
    last = max((i for i in range(width) if not is_empty(header[i]) or any(not is_empty(r[i]) for r in grid[h + 1:])), default=-1)
    columns = unique_names([clean_spaces(str(v)) if not is_empty(v) else f"Колонка {i + 1}" for i, v in enumerate(header[: last + 1])])
    rows, numbers = [], []
    for idx, r in enumerate(grid[h + 1:], start=h + 2):
        rows.append([_cell(v) for v in r[: last + 1]])
        numbers.append(idx)
    return Table(columns, rows, name, numbers)


def _cell(v: Any) -> Any:
    if isinstance(v, str):
        return v if v.strip() else None
    return v


def unique_names(names: Iterable[str]) -> list[str]:
    seen: dict[str, int] = {}
    out = []
    for n in names:
        if n in seen:
            seen[n] += 1
            out.append(f"{n} ({seen[n]})")
        else:
            seen[n] = 1
            out.append(n)
    return out


# ================================================================== запись
HEADER_FILL = PatternFill("solid", fgColor="1F4E78")
HEADER_FONT = Font(bold=True, color="FFFFFF")
CHANGED_FILL = PatternFill("solid", fgColor="FFF2CC")   # изменено программой
PROBLEM_FILL = PatternFill("solid", fgColor="F8CBAD")   # не распознано — проверить
ADDED_FILL = PatternFill("solid", fgColor="E2EFDA")     # дополнено из дубля
THIN = Side(style="thin", color="BFBFBF")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)


@dataclass
class Sheet:
    title: str
    columns: list[str]
    rows: list[list[Any]]
    fills: dict[tuple[int, int], PatternFill] = field(default_factory=dict)  # (строка, колонка) с 0
    plain: bool = False  # лист-отчёт: без шапки и сетки


def safe_sheet_title(title: str, used: set[str]) -> str:
    t = re.sub(r"[\[\]:*?/\\]", "_", title).strip() or "Лист"
    t = t[:31]
    base, n = t, 2
    while t in used:
        suffix = f" ({n})"
        t = base[: 31 - len(suffix)] + suffix
        n += 1
    used.add(t)
    return t


def write_xlsx(path: Path, sheets: list[Sheet]) -> Path:
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    wb = Workbook()
    wb.remove(wb.active)
    used: set[str] = set()
    for sh in sheets:
        ws = wb.create_sheet(safe_sheet_title(sh.title, used))
        if sh.plain:
            _write_plain(ws, sh.rows)
            continue
        ws.append(sh.columns)
        for c in range(1, len(sh.columns) + 1):
            cell = ws.cell(1, c)
            cell.fill, cell.font, cell.border = HEADER_FILL, HEADER_FONT, BORDER
            cell.alignment = Alignment(wrap_text=True, vertical="center")
        for r_idx, row in enumerate(sh.rows):
            ws.append([_excel_value(v) for v in row])
            for c_idx, v in enumerate(row):
                cell = ws.cell(r_idx + 2, c_idx + 1)
                cell.border = BORDER
                if isinstance(v, str) and v.startswith("="):
                    cell.data_type = "s"  # текст «=…» из исходника не должен стать формулой
                fmt = _number_format(v)
                if fmt:
                    cell.number_format = fmt
                fill = sh.fills.get((r_idx, c_idx))
                if fill:
                    cell.fill = fill
        ws.freeze_panes = "A2"
        if sh.rows:
            ws.auto_filter.ref = f"A1:{get_column_letter(len(sh.columns))}{len(sh.rows) + 1}"
        _autowidth(ws, sh.columns, sh.rows)
    try:
        wb.save(path)
    except PermissionError as exc:
        raise DocToolsError(f"Не удалось записать {path.name}: файл открыт в Excel? Закройте его и повторите.") from exc
    return path


def _write_plain(ws, rows: list[list[Any]]) -> None:
    for row in rows:
        ws.append([_excel_value(v) for v in row])
    if rows:
        ws.cell(1, 1).font = Font(bold=True, size=14)
    ws.column_dimensions["A"].width = 60
    for col in "BCDEFG":
        ws.column_dimensions[col].width = 22


def _excel_value(v: Any) -> Any:
    if isinstance(v, (list, dict)):
        return json.dumps(v, ensure_ascii=False)
    return v


def _has_time(v: datetime) -> bool:
    """Excel хранит даты как datetime в полночь — время показываем, только если оно есть."""
    return (v.hour, v.minute, v.second) != (0, 0, 0)


def _number_format(v: Any) -> str | None:
    if isinstance(v, bool):
        return None
    if isinstance(v, datetime):
        return "DD.MM.YYYY HH:MM" if _has_time(v) else "DD.MM.YYYY"
    if isinstance(v, date):
        return "DD.MM.YYYY"
    if isinstance(v, int):
        return "#,##0"
    if isinstance(v, float):
        return "#,##0.00"
    return None


def _autowidth(ws, columns: list[str], rows: list[list[Any]]) -> None:
    for i, name in enumerate(columns):
        longest = len(str(name))
        for row in rows[:2000]:
            if i < len(row) and row[i] is not None:
                v = row[i]
                n = 10 if isinstance(v, (date, datetime)) else len(str(v))
                longest = max(longest, n)
        width = min(max(longest + 2, 8), 60)
        ws.column_dimensions[get_column_letter(i + 1)].width = width
        if longest > 60:
            for cell in ws[get_column_letter(i + 1)][1:]:
                cell.alignment = Alignment(wrap_text=True, vertical="top")


def to_text(v: Any) -> str:
    """Значение для CSV/отчётов: даты — ДД.ММ.ГГГГ, дробные — с запятой."""
    if v is None:
        return ""
    if isinstance(v, datetime) and _has_time(v):
        return v.strftime("%d.%m.%Y %H:%M")
    if isinstance(v, date):
        return v.strftime("%d.%m.%Y")
    if isinstance(v, float):
        return (f"{v:.2f}" if not v.is_integer() else str(int(v))).replace(".", ",")
    return str(v)


def write_csv(path: Path, table: Table, sep: str = ";", encoding: str = "utf-8-sig") -> Path:
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", newline="", encoding=encoding, errors="replace") as fh:
        w = csv.writer(fh, delimiter=sep)
        w.writerow(table.columns)
        for row in table.rows:
            w.writerow([to_text(v) for v in row])
    return path


def json_value(v: Any) -> Any:
    if isinstance(v, datetime):
        return v.isoformat(timespec="minutes") if _has_time(v) else v.date().isoformat()
    if isinstance(v, date):
        return v.isoformat()
    return v


def table_records(table: Table) -> list[dict[str, Any]]:
    return [{c: json_value(v) for c, v in zip(table.columns, row)} for row in table.rows]


def write_json(path: Path, tables: list[Table]) -> Path:
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    if len(tables) == 1:
        data: Any = table_records(tables[0])
    else:
        data = {t.name.split(" / ")[-1]: table_records(t) for t in tables}
    path.write_text(json.dumps(data, ensure_ascii=False, indent=2), encoding="utf-8")
    return path


# ============================================== типизация колонок (общая)
_CODE_HINTS = ("артикул", "код", "инн", "кпп", "огрн", "бик", "счет", "счёт", "телефон", "тел",
               "phone", "индекс", "№", "номер", "паспорт", "снилс", "sku", "id")


def is_code_column(name: str) -> bool:
    n = name.casefold()
    return any(h in n for h in _CODE_HINTS) and not any(w in n for w in ("сумм", "цен", "кол"))


@dataclass
class TypingResult:
    kind: dict[str, str] = field(default_factory=dict)        # колонка -> «число»/«дата»/«текст»
    failed: list[tuple[int, str, Any]] = field(default_factory=list)  # (строка, колонка, значение)


def infer_and_convert(table: Table, share: float = 0.6) -> TypingResult:
    """Приводит строки к числам и датам по колонкам. Колонка считается числовой
    (датовой), если так распознаётся не меньше share непустых значений; остальные
    значения такой колонки остаются текстом и попадают в failed."""
    res = TypingResult()
    for ci, name in enumerate(table.columns):
        values = [(ri, row[ci]) for ri, row in enumerate(table.rows) if ci < len(row) and not is_empty(row[ci])]
        if not values or is_code_column(name):
            res.kind[name] = "текст"
            continue
        dates = {ri: parse_date(v) for ri, v in values if isinstance(v, str) or isinstance(v, (date, datetime))}
        dates = {k: d for k, d in dates.items() if d is not None}
        nums = {ri: parse_number(v) for ri, v in values if not isinstance(v, (date, datetime))}
        nums = {k: n for k, n in nums.items() if n is not None}
        if len(dates) >= share * len(values):
            kind, parsed = "дата", dates
        elif len(nums) >= share * len(values):
            kind, parsed = "число", nums
        else:
            res.kind[name] = "текст"
            continue
        res.kind[name] = kind
        for ri, v in values:
            if ri in parsed:
                table.rows[ri][ci] = parsed[ri]
            else:
                res.failed.append((ri, name, v))
    return res
