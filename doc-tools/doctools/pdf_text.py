"""pdf2docx / pdftext: текст PDF в Word или TXT с абзацами, заголовками, списками и таблицами.

Как восстанавливается структура (без OCR, только по текстовому слою):
- строки собираются из слов по вертикали, абзацы — по отступам, интервалам
  и коротким последним строкам;
- заголовки — по размеру шрифта (крупнее основного) или по жирной короткой строке;
- повторяющиеся колонтитулы и номера страниц удаляются;
- переносы по слогам («распо-/знавание») склеиваются, абзац, разорванный
  между страницами, собирается обратно;
- таблицы с сеткой переносятся таблицами Word (в TXT — через табуляцию);
- страницы-сканы помечаются «нужен OCR».
"""

from __future__ import annotations

import re
from collections import Counter
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any

from docx import Document
from docx.enum.text import WD_COLOR_INDEX
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor

from . import DocToolsError
from .normalize import clean_spaces
from .pdf_tables import find_page_tables, open_pdf, page_has_text, parse_pages
from .report import Report

BULLET = re.compile(r"^([•●▪■◦\-–—*])\s+")
NUMBERED = re.compile(r"^(\d{1,2}[.)]|[а-яa-z][.)])\s+", re.I)
SENT_END = tuple(".!?:;»\")")
LABEL = re.compile(r"^[А-ЯЁA-Z][\w .№-]{0,30}:\s")  # «Период: …», «ИНН: …» — реквизиты с новой строки


@dataclass
class Line:
    text: str
    top: float
    bottom: float
    x0: float
    x1: float
    size: float
    bold: bool

    @property
    def height(self) -> float:
        return self.bottom - self.top


@dataclass
class Block:
    kind: str                      # heading | para | bullet | table | ocr
    text: str = ""
    level: int = 0
    rows: list[list[str | None]] = field(default_factory=list)
    page: int = 0
    size: float = 0.0
    bold: bool = False
    top: float = 0.0


@dataclass
class Stats:
    removed_running: Counter = field(default_factory=Counter)
    hyphen_joins: int = 0
    page_joins: int = 0
    ocr_pages: list[int] = field(default_factory=list)


# ------------------------------------------------------------ строки страницы
def _page_lines(page: Any, exclude: list[tuple[float, float, float, float]]) -> list[Line]:
    words = page.extract_words(extra_attrs=["size", "fontname"], keep_blank_chars=False, x_tolerance=2)

    def inside(w: dict) -> bool:
        cx, cy = (w["x0"] + w["x1"]) / 2, (w["top"] + w["bottom"]) / 2
        return any(x0 - 1 <= cx <= x1 + 1 and t - 1 <= cy <= b + 1 for x0, t, x1, b in exclude)

    words = [w for w in words if not inside(w)]
    words.sort(key=lambda w: (round(w["top"]), w["x0"]))
    groups: list[list[dict]] = []
    for w in words:
        if groups and abs(w["top"] - groups[-1][0]["top"]) <= max(2.0, w["size"] * 0.3):
            groups[-1].append(w)
        else:
            groups.append([w])
    lines = []
    for g in groups:
        g.sort(key=lambda w: w["x0"])
        text = clean_spaces(" ".join(w["text"] for w in g))
        if not text:
            continue
        sizes = sorted(w["size"] for w in g)
        bold = sum("bold" in w["fontname"].lower() for w in g) >= 0.8 * len(g)
        lines.append(Line(text, min(w["top"] for w in g), max(w["bottom"] for w in g),
                          g[0]["x0"], g[-1]["x1"], round(sizes[len(sizes) // 2], 1), bold))
    return lines


def _running_key(text: str) -> str:
    return re.sub(r"\d+", "#", text.casefold()).strip()


def _find_running(pages_lines: list[tuple[float, list[Line]]]) -> set[str]:
    """Колонтитулы: одинаковые (с точностью до цифр) строки у края листа на половине страниц и больше."""
    if len(pages_lines) < 2:
        return set()
    cnt: Counter = Counter()
    for height, lines in pages_lines:
        keys = {_running_key(ln.text) for ln in lines if ln.top < height * 0.08 or ln.bottom > height * 0.92}
        cnt.update(keys)
    need = max(2, len(pages_lines) // 2)
    return {k for k, n in cnt.items() if n >= need}


def _is_running(ln: Line, height: float, running: set[str]) -> bool:
    edge = ln.top < height * 0.08 or ln.bottom > height * 0.92
    return edge and (_running_key(ln.text) in running or bool(re.fullmatch(r"[-–— ]*\d+[-–— ]*", ln.text)))


# ------------------------------------------------------------ сборка блоков
def extract_blocks(src: Path, pages: str | None = None, tables: bool = True) -> tuple[list[Block], Stats, int]:
    stats = Stats()
    blocks: list[Block] = []
    with open_pdf(src) as pdf:
        total = len(pdf.pages)
        numbers = parse_pages(pages, total)
        per_page: list[tuple[int, float, float, list[Line], list[Block]]] = []
        hint: list[float] | None = None
        for pno in numbers:
            page = pdf.pages[pno - 1]
            if not page_has_text(page):
                if page.images:
                    stats.ocr_pages.append(pno)
                per_page.append((pno, float(page.height), float(page.width), [], []))
                continue
            tbl_blocks: list[Block] = []
            boxes = []
            if tables:
                found = find_page_tables(page, pno, "auto", hint)
                hint = found[-1].xs if found else None
                for t in found:
                    tbl_blocks.append(Block("table", rows=t.rows, page=pno, top=t.top))
                    boxes.append(t.bbox)
            per_page.append((pno, float(page.height), float(page.width), _page_lines(page, boxes), tbl_blocks))

    running = _find_running([(h, ls) for _, h, _, ls, _ in per_page if ls])
    all_lines = [ln for _, h, _, ls, _ in per_page for ln in ls if not _is_running(ln, h, running)]
    if not all_lines and not any(tb for *_, tb in per_page) and not stats.ocr_pages:
        raise DocToolsError(f"{Path(src).name}: в PDF не найден текст")
    body = _body_size(all_lines)
    heading_sizes = sorted({ln.size for ln in all_lines if ln.size >= body * 1.15}, reverse=True)

    for pno, height, width, lines, tbl_blocks in per_page:
        if pno in stats.ocr_pages:
            blocks.append(Block("ocr", page=pno))
            continue
        kept = []
        for ln in lines:
            if _is_running(ln, height, running):
                stats.removed_running[ln.text] += 1
                continue
            kept.append(ln)
        page_blocks = _group(kept, body, heading_sizes, pno, stats) + tbl_blocks
        page_blocks.sort(key=lambda b: b.top)
        _join_across_pages(blocks, page_blocks, stats)
        blocks += page_blocks
    for b in blocks:
        if b.kind == "heading" and len(b.text) > 150:
            b.kind = "para"
    return _merge_tables(blocks), stats, total


def _body_size(lines: list[Line]) -> float:
    c: Counter = Counter()
    for ln in lines:
        c[ln.size] += len(ln.text)
    return c.most_common(1)[0][0] if c else 10.0


def _is_heading(ln: Line, body: float) -> bool:
    if ln.size >= body * 1.15:
        return True
    return ln.bold and len(ln.text) <= 90 and not ln.text.endswith((".", ",", ";")) and not BULLET.match(ln.text)


def _group(lines: list[Line], body: float, heading_sizes: list[float], pno: int, stats: Stats) -> list[Block]:
    if not lines:
        return []
    left = min(ln.x0 for ln in lines)
    right = max(ln.x1 for ln in lines)
    width = max(right - left, 1)
    out: list[Block] = []
    prev: Line | None = None
    for ln in lines:
        heading = _is_heading(ln, body)
        marker = BULLET.match(ln.text) or NUMBERED.match(ln.text) or LABEL.match(ln.text)
        cur = out[-1] if out else None
        new = True
        if cur is not None and prev is not None and not marker:
            gap = ln.top - prev.bottom
            same_kind = (cur.kind == "heading") == heading
            close = gap <= max(prev.height * 0.6, 3)
            # строка внутри абзаца доходит почти до правого края; короткая — значит, абзац кончился
            short_prev = prev.x1 < right - width * 0.15
            indented = ln.x0 > left + 12 and cur.kind == "para"
            hanging = cur.kind == "bullet" and ln.x0 > left + 4
            if same_kind and close and abs(ln.size - prev.size) < 0.6:
                if cur.kind == "heading":
                    new = False
                elif cur.kind == "bullet":
                    new = not hanging and short_prev
                else:
                    new = short_prev or indented
        if new:
            if heading:
                level = heading_sizes.index(ln.size) + 1 if ln.size in heading_sizes else len(heading_sizes) + 1
                out.append(Block("heading", ln.text, min(level, 3), page=pno, size=ln.size, bold=ln.bold, top=ln.top))
            elif BULLET.match(ln.text):
                out.append(Block("bullet", BULLET.sub("", ln.text), page=pno, size=ln.size, top=ln.top))
            else:
                out.append(Block("para", ln.text, page=pno, size=ln.size, top=ln.top))
        else:
            assert cur is not None
            cur.text = _join_text(cur.text, ln.text, stats)
        prev = ln
    return out


def _join_text(a: str, b: str, stats: Stats) -> str:
    if re.search(r"\w-$", a) and b[:1].islower():
        stats.hyphen_joins += 1
        return a[:-1] + b
    return a + " " + b


def _join_across_pages(done: list[Block], page_blocks: list[Block], stats: Stats) -> None:
    if not done or not page_blocks:
        return
    last, first = done[-1], page_blocks[0]
    if last.kind == "para" and first.kind == "para" and not last.text.endswith(SENT_END) and first.text[:1].islower():
        last.text = _join_text(last.text, first.text, stats)
        page_blocks.pop(0)
        stats.page_joins += 1


def _merge_tables(blocks: list[Block]) -> list[Block]:
    """Таблица, продолженная на следующей странице, — одна таблица; повтор шапки убираем."""
    out: list[Block] = []
    for b in blocks:
        prev = out[-1] if out else None
        if (b.kind == "table" and prev is not None and prev.kind == "table" and b.page == prev.page + 1
                and len(b.rows[0]) == len(prev.rows[0])):
            rows = b.rows[1:] if b.rows[0] == prev.rows[0] else b.rows
            prev.rows += rows
            prev.page = b.page
            continue
        out.append(b)
    return out


# ------------------------------------------------------------ запись DOCX/TXT
def _repeat_header(row: Any) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    el = OxmlElement("w:tblHeader")
    el.set(qn("w:val"), "true")
    tr_pr.append(el)


def write_docx(blocks: list[Block], dst: Path, title: str | None = None) -> Path:
    doc = Document()
    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(11)
    normal.paragraph_format.space_after = Pt(6)
    for b in blocks:
        if b.kind == "heading":
            doc.add_heading(b.text, level=b.level)
        elif b.kind == "bullet":
            doc.add_paragraph(b.text, style="List Bullet")
        elif b.kind == "para":
            doc.add_paragraph(b.text)
        elif b.kind == "table":
            ncols = max(len(r) for r in b.rows)
            table = doc.add_table(rows=0, cols=ncols)
            table.style = "Table Grid"
            for i, r in enumerate(b.rows):
                cells = table.add_row().cells
                for j in range(ncols):
                    val = r[j] if j < len(r) else None
                    cells[j].text = val or ""
                    for p in cells[j].paragraphs:
                        p.paragraph_format.space_after = Pt(0)
                        for run in p.runs:
                            run.font.size = Pt(9)
                            run.bold = i == 0
                if i == 0:
                    _repeat_header(table.rows[0])
            doc.add_paragraph()
        elif b.kind == "ocr":
            p = doc.add_paragraph()
            run = p.add_run(f"[Страница {b.page}: скан без текстового слоя — текст не извлечён, нужен OCR]")
            run.italic = True
            run.font.color.rgb = RGBColor(0x99, 0x33, 0x00)
            run.font.highlight_color = WD_COLOR_INDEX.YELLOW
    if title:
        doc.core_properties.title = title
    dst.parent.mkdir(parents=True, exist_ok=True)
    try:
        doc.save(dst)
    except PermissionError as exc:
        raise DocToolsError(f"Не удалось записать {dst.name}: файл открыт в Word? Закройте его и повторите.") from exc
    return dst


def blocks_to_text(blocks: list[Block]) -> str:
    parts: list[str] = []
    for b in blocks:
        if b.kind == "heading":
            under = "=" if b.level == 1 else "-"
            parts.append(b.text + "\n" + under * min(len(b.text), 80))
        elif b.kind == "bullet":
            parts.append("• " + b.text)
        elif b.kind == "para":
            parts.append(b.text)
        elif b.kind == "table":
            parts.append("\n".join("\t".join(c or "" for c in r) for r in b.rows))
        elif b.kind == "ocr":
            parts.append(f"[Страница {b.page}: скан без текстового слоя — нужен OCR]")
    return "\n\n".join(parts) + "\n"


def _report(src: Path, blocks: list[Block], stats: Stats, total: int, what: str) -> Report:
    rep = Report(f"PDF -> {what}: {Path(src).name}")
    kinds = Counter(b.kind for b in blocks)
    rep.stat("Страниц в PDF", total)
    rep.stat("Заголовков", kinds["heading"])
    rep.stat("Абзацев", kinds["para"])
    rep.stat("Пунктов списков", kinds["bullet"])
    rep.stat("Таблиц", kinds["table"])
    rep.stat("Страниц «нужен OCR»", len(stats.ocr_pages))
    if stats.removed_running:
        sample = "; ".join(f"«{t}»" for t in list(stats.removed_running)[:3])
        rep.info(f"Удалены колонтитулы и номера страниц: {sum(stats.removed_running.values())} строк (например, {sample})")
    if stats.hyphen_joins:
        rep.info(f"Склеено переносов по слогам: {stats.hyphen_joins}")
    if stats.page_joins:
        rep.info(f"Абзацев, собранных через границу страниц: {stats.page_joins}")
    if stats.ocr_pages:
        rep.warn(f"Страницы-сканы без текстового слоя: {', '.join(map(str, stats.ocr_pages))} — "
                 "на их месте стоит пометка «нужен OCR» (распознавание — отдельно, по согласованию)")
    rep.info("Заголовки определены по размеру и жирности шрифта; сложная вёрстка (колонки, врезки) — проверить глазами")
    return rep


def pdf_to_docx(src: Path, dst: Path | None = None, *, pages: str | None = None, tables: bool = True) -> tuple[Path, Report]:
    src = Path(src)
    dst = Path(dst) if dst else src.with_suffix(".docx")
    blocks, stats, total = extract_blocks(src, pages, tables)
    write_docx(blocks, dst, title=src.stem)
    rep = _report(src, blocks, stats, total, "Word")
    rep.output(dst)
    return dst, rep


def pdf_to_text(src: Path, dst: Path | None = None, *, pages: str | None = None, tables: bool = True) -> tuple[Path, Report]:
    src = Path(src)
    dst = Path(dst) if dst else src.with_suffix(".txt")
    blocks, stats, total = extract_blocks(src, pages, tables)
    dst.parent.mkdir(parents=True, exist_ok=True)
    dst.write_text(blocks_to_text(blocks), encoding="utf-8")
    rep = _report(src, blocks, stats, total, "TXT")
    rep.output(dst)
    return dst, rep
