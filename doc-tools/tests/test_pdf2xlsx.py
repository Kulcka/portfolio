from datetime import date

from openpyxl import load_workbook

from doctools.cli import main
from doctools.pdf_tables import pdf_to_xlsx
from doctools.samples import (DISCOUNTS, PRICE_ITEMS, PRICE_ON_REQUEST, PRICE_SECTIONS, STATEMENT_OPS,
                              statement_operations)


def _rows(ws):
    return list(ws.iter_rows(values_only=True))


def test_price_table_joined_across_pages(samples, out):
    dst, rep = pdf_to_xlsx(samples.price_pdf, out / "price.xlsx")
    wb = load_workbook(dst)
    assert wb.sheetnames == ["Таблица 1 (стр. 1-2)", "Таблица 2 (стр. 3)", "Отчёт"]
    rows = _rows(wb.worksheets[0])
    assert rows[0] == ("Раздел", "Артикул", "Наименование", "Ед.", "Цена, ₽", "Остаток", "Поступление")
    body = rows[1:]
    assert len(body) == PRICE_ITEMS                      # повтор шапки на стр. 2 убран, разделы — не строки
    assert [r[1] for r in body] == [f"СС-{i:04d}" for i in range(1, PRICE_ITEMS + 1)]
    assert {r[0] for r in body} == set(PRICE_SECTIONS)   # колонка «Раздел» заполнена
    prices = [r[4] for r in body]
    assert sum(isinstance(p, (int, float)) for p in prices) == PRICE_ITEMS - len(PRICE_ON_REQUEST)
    assert all(isinstance(r[6], date) for r in body)     # даты стали датами Excel
    assert any(r[2] == "Гайка М10 самоконтрящаяся с нейлоновым кольцом, DIN 985" for r in body)  # перенос склеен
    assert any("Не распознано как число/дата значений: 4" in w for w in rep.warnings)


def test_discount_table(samples, out):
    dst, _ = pdf_to_xlsx(samples.price_pdf, out / "price.xlsx")
    rows = _rows(load_workbook(dst).worksheets[1])
    assert rows[0] == ("Сумма заказа от, ₽", "Скидка, %")
    assert [r[1] for r in rows[1:]] == [int(b) for _, b in DISCOUNTS]
    assert rows[-1][0] == 1_000_000


def test_statement_without_vertical_lines(samples, out):
    dst, rep = pdf_to_xlsx(samples.statement_pdf, out / "st.xlsx")
    wb = load_workbook(dst)
    assert len(wb.sheetnames) == 2
    rows = _rows(wb.worksheets[0])
    assert rows[0] == ("Дата", "№ док.", "Контрагент", "Назначение платежа", "Поступление", "Списание")
    ops = statement_operations()
    body, total = rows[1:-1], rows[-1]
    assert len(body) == STATEMENT_OPS
    assert total[3] == "Обороты за период"
    assert round(sum(r[4] or 0 for r in body), 2) == round(sum(o[4] for o in ops), 2) == total[4]
    assert round(sum(r[5] or 0 for r in body), 2) == round(sum(o[5] for o in ops), 2) == total[5]
    assert [r[3] for r in body] == [o[3] for o in ops]   # длинное «Назначение» не разорвано по колонкам
    assert any("продолжение без шапки" in t for _, t in rep.items)


def test_scan_marked_ocr(samples, out):
    _, rep = pdf_to_xlsx(samples.scan_pdf, out / "scan.xlsx")
    assert rep.stats["Страниц «нужен OCR»"] == 1
    assert any("нужен OCR" in w for w in rep.warnings)


def test_cli_pdf2xlsx_and_report(samples, out, capsys):
    code = main(["pdf2xlsx", str(samples.price_pdf), "-o", str(out / "p.xlsx"), "--pages", "3", "--source-page"])
    assert code == 0
    assert (out / "p_отчёт.md").exists()
    rows = _rows(load_workbook(out / "p.xlsx").worksheets[0])
    assert rows[0][-1] == "Стр. PDF" and rows[1][-1] == 3
    assert "Таблиц найдено: 1" in capsys.readouterr().out


def test_cli_errors(out, capsys):
    assert main(["pdf2xlsx", str(out / "нет.pdf")]) == 2
    bad = out / "bad.pdf"
    bad.write_bytes(b"not a pdf")
    assert main(["pdf2xlsx", str(bad)]) == 2
    assert main(["pdf2xlsx", str(bad), "--pages", "x"]) == 2
    assert "Ошибка" in capsys.readouterr().err
