from datetime import date

import pytest
from openpyxl import load_workbook

from doctools import DocToolsError
from doctools.clean import clean_file, detect_type
from doctools.cli import main
from doctools.merge import ColumnConfig, merge_files


def _sheet(path, name):
    rows = list(load_workbook(path)[name].iter_rows(values_only=True))
    return rows[0], rows[1:]


@pytest.fixture()
def merged(samples, out):
    dst, rep = merge_files([samples.sales_xlsx, samples.crm_xlsx, samples.site_csv], out / "merged.xlsx")
    return dst, rep


def test_merge_columns_and_sources(merged):
    dst, rep = merged
    header, rows = _sheet(dst, "Объединено")
    assert header == ("ФИО", "Телефон", "Email", "Город", "ИНН", "Дата обращения", "Сумма заказа",
                      "Комментарий", "Источник", "Строка источника")
    assert len(rows) == 9 + 6 + 4                         # пустая строка отдела продаж пропущена
    sources = {r[8] for r in rows}
    assert sources == {"клиенты_отдел_продаж.xlsx", "выгрузка_CRM.xlsx", "заявки_с_сайта.csv"}
    crm_first = next(r for r in rows if r[8] == "выгрузка_CRM.xlsx")
    assert crm_first[0] == "Смирнова Ольга" and crm_first[9] == 4   # шапка CRM на 3-й строке найдена сама
    assert any("нет колонок ИНН" in w for w in rep.warnings)          # у CSV нет ИНН
    assert any("Комментарий" in w for w in rep.warnings)


def test_merge_fuzzy_header_and_config(tmp_path):
    cfg = tmp_path / "c.yaml"
    cfg.write_text("columns:\n  Телефон: [телефон]\n  Город: [город]\nfuzzy_threshold: 85\n", encoding="utf-8")
    c = ColumnConfig.load(cfg)
    assert c.match("Телефонн")[0] == "Телефон" and c.match("Телефонн")[1].startswith("похоже")
    assert c.match("Скидка") == (None, "нет в словаре")
    bad = tmp_path / "bad.yaml"
    bad.write_text("foo: 1\n", encoding="utf-8")
    with pytest.raises(DocToolsError):
        ColumnConfig.load(bad)


def test_detect_types():
    assert detect_type("Номер телефона") == "phone"
    assert detect_type("Тел.") == "phone"
    assert detect_type("ИНН организации") == "inn"
    assert detect_type("E-mail") == "email"
    assert detect_type("Дата обращения") == "date"
    assert detect_type("Сумма, ₽") == "number"
    assert detect_type("Контактное лицо") == "name"
    assert detect_type("Комментарий") == "text"


def test_clean_after_merge(merged, out):
    dst, rep = clean_file(merged[0], out / "clean.xlsx", dedup_keys=["ФИО", "Телефон"], fuzzy_keys=["ФИО"])
    header, rows = _sheet(dst, "Результат")
    assert len(rows) == 16                                                # 19 - 3 точных дубля
    by_name = {r[0]: r for r in rows}
    petrov = by_name["Петров Иван"]                                       # «  петров  иван » -> регистр и пробелы
    assert petrov[1] == "+7 (900) 000-23-45" and petrov[2] == "petrov@example.com"
    assert isinstance(petrov[5], date) and petrov[6] == 8300
    ivanova = by_name["Иванова Анна Сергеевна"]
    assert ivanova[4] == "0000010016"                                     # ИНН: восстановлены ведущие нули
    assert ivanova[7] == "просила счёт на e-mail"                         # дополнено из дубля
    assert all(isinstance(r[6], (int, float)) for r in rows)             # все суммы — числа
    assert rep.stats["Точных дублей удалено"] == 3
    assert rep.stats["Групп похожих строк (проверить)"] == 2              # Смирнова О. / Кузнецов с опечаткой
    reasons = {r[3] for r in rep.tables[0].rows}                          # таблица «Не распознано»
    assert reasons == {"контрольные цифры не сходятся", "телефон не распознан", "некорректный e-mail"}
    _, changes = _sheet(dst, "Изменения")
    assert len(changes) == rep.stats["Ячеек изменено"] > 30
    _, removed = _sheet(dst, "Удалённые строки")
    assert sorted(r[1] for r in removed) == ["дубль"] * 3
    _, before = _sheet(dst, "До")
    assert len(before) == 19


def test_clean_drop_fuzzy_and_csv(merged, out):
    assert main(["clean", str(merged[0]), "-o", str(out / "c.xlsx"), "--dedup-keys", "ФИО,Телефон",
                 "--fuzzy-keys", "ФИО", "--drop-fuzzy"]) == 0
    _, rows = _sheet(out / "c.xlsx", "Результат")
    assert len(rows) == 14
    assert main(["clean", str(merged[0]), "-o", str(out / "c.csv")]) == 0
    text = (out / "c.csv").read_text(encoding="utf-8-sig")
    assert text.splitlines()[0].startswith("ФИО;Телефон;")


def test_clean_type_override_and_errors(samples, out, capsys):
    assert main(["clean", str(samples.sales_xlsx), "-o", str(out / "s.xlsx"), "--type", "Город=текст"]) == 0
    _, rows = _sheet(out / "s.xlsx", "Результат")
    assert rows[1][3] == "демоград"                                       # город оставлен как есть
    assert len(rows) == 8                                                 # пустая строка и точный дубль удалены
    assert main(["clean", str(samples.sales_xlsx), "--type", "Нет такой=телефон"]) == 2
    assert main(["clean", str(samples.sales_xlsx), "--type", "Город=космос"]) == 2
    assert "Нет колонки" in capsys.readouterr().err
