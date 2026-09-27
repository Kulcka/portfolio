from docx import Document

from doctools.cli import main
from doctools.pdf_text import pdf_to_docx, pdf_to_text
from doctools.samples import PRICE_ITEMS, PRICE_SECTIONS, STATEMENT_OPS


def test_pdf2docx_structure(samples, out):
    dst, rep = pdf_to_docx(samples.price_pdf, out / "price.docx")
    doc = Document(dst)
    heads = [(p.style.name, p.text) for p in doc.paragraphs if p.style.name.startswith("Heading")]
    assert heads[0] == ("Heading 1", "Прайс-лист")
    assert ("Heading 2", "Условия поставки") in heads
    bullets = [p.text for p in doc.paragraphs if p.style.name == "List Bullet"]
    assert len(bullets) == 3 and bullets[0].startswith("Доставка по Демограду")
    text = "\n".join(p.text for p in doc.paragraphs)
    assert "Стр. 1 из 3" not in text and "info@example.com" not in text   # колонтитулы убраны
    intro = next(p.text for p in doc.paragraphs if p.text.startswith("Уважаемые партнёры"))
    assert intro.endswith("уточняйте у менеджера.")                        # абзац собран из строк целиком
    assert len(doc.tables) == 2
    assert len(doc.tables[0].rows) == 1 + PRICE_ITEMS + len(PRICE_SECTIONS)  # склеено со стр. 2, шапка одна
    assert rep.stats["Таблиц"] == 2


def test_pdftext_statement(samples, out):
    dst, _ = pdf_to_text(samples.statement_pdf, out / "st.txt")
    lines = dst.read_text(encoding="utf-8").splitlines()
    assert "Период: 01.09.2026 — 15.09.2026" in lines
    table_lines = [ln for ln in lines if ln.count("\t") == 5]
    assert len(table_lines) == 1 + STATEMENT_OPS + 1                       # шапка + операции + обороты


def test_scan_placeholder(samples, out):
    assert main(["pdf2docx", str(samples.scan_pdf), "-o", str(out / "scan.docx")]) == 0
    assert "нужен OCR" in Document(out / "scan.docx").paragraphs[0].text
    assert main(["pdftext", str(samples.scan_pdf), "-o", str(out / "scan.txt")]) == 0
    assert "нужен OCR" in (out / "scan.txt").read_text(encoding="utf-8")
