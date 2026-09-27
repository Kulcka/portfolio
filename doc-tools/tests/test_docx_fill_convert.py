import json

import pytest
from docx import Document
from openpyxl import Workbook, load_workbook

from doctools import DocToolsError
from doctools.cli import main
from doctools.convert import convert_files
from doctools.docx_fill import fill_documents
from doctools.samples import RECIPIENTS


def _all_text(doc):
    parts = [p.text for p in doc.paragraphs]
    parts += [c.text for t in doc.tables for r in t.rows for c in r.cells]
    for s in doc.sections:
        parts += [p.text for p in s.header.paragraphs] + [p.text for p in s.footer.paragraphs]
    return "\n".join(parts)


def test_fill_letters(samples, out):
    made, rep = fill_documents(samples.letter_docx, samples.recipients_xlsx, out / "letters",
                               name_pattern="Письмо_{{Номер договора}}_{{Компания}}")
    assert len(made) == len(RECIPIENTS)
    assert made[0].name == "Письмо_17-Д_ООО «Ромашка-Демо».docx"
    text = _all_text(Document(made[0]))
    assert "{{" not in text and "}}" not in text
    assert "О продлении договора № 17-Д" in text                          # плейсхолдер, разбитый на 3 фрагмента
    assert "Уважаемая Анна Сергеевна!" in text
    assert "Иванова А. С." in text                                        # |инициалы
    assert "от 01.10.2025" in text                                        # |дата
    assert "1 250 000,00" in text                               # |деньги
    assert "Один миллион двести пятьдесят тысяч рублей 00 копеек" in text  # |прописью
    assert "Письмо подготовлено для ООО «Ромашка-Демо»" in text           # колонтитул
    # у Кузнецова нет должности — документ создан, но попал в отчёт
    assert rep.tables and rep.tables[0].rows[0][2] == "Должность"


def test_fill_missing_column(samples, out, tmp_path):
    wb = Workbook()
    wb.active.append(["ФИО"])
    wb.active.append(["Иванова Анна"])
    data = tmp_path / "d.xlsx"
    wb.save(data)
    with pytest.raises(DocToolsError, match="нет колонок для плейсхолдеров"):
        fill_documents(samples.letter_docx, data, out / "x")
    made, rep = fill_documents(samples.letter_docx, data, out / "x", allow_missing=True)
    assert len(made) == 1 and rep.warnings


def test_convert_roundtrip(samples, out):
    made, _ = convert_files([samples.site_csv], "xlsx", out)
    rows = list(load_workbook(made[0]).active.iter_rows(values_only=True))
    assert rows[1][5] == 3200 and rows[1][1] == "8 900 000 78 90"         # число стало числом, телефон — текстом
    made, _ = convert_files([made[0]], "json", out / "j")
    data = json.loads(made[0].read_text(encoding="utf-8"))
    assert data[0]["Имя"] == "Федоров Павел" and data[0]["Дата заявки"] == "2026-09-10"
    made, _ = convert_files([made[0]], "csv", out / "c", sep=";")
    assert made[0].read_text(encoding="utf-8-sig").splitlines()[1].startswith("Федоров Павел;")


def test_convert_multisheet_and_skips(out, capsys):
    wb = Workbook()
    wb.active.title = "Январь"
    wb.active.append(["a", "b"]); wb.active.append([1, 2])
    ws = wb.create_sheet("Февраль")
    ws.append(["a", "b"]); ws.append([3, 4])
    src = out / "двалиста.xlsx"
    wb.save(src)
    made, rep = convert_files([src], "csv", out / "csv")
    assert sorted(p.name for p in made) == ["двалиста_Февраль.csv", "двалиста_Январь.csv"]
    made, rep = convert_files([src], "json", out / "json")
    assert set(json.loads(made[0].read_text(encoding="utf-8"))) == {"Январь", "Февраль"}
    assert main(["convert", str(src), "--to", "xlsx", "-o", str(out / "same")]) == 0   # уже xlsx — пропуск
    assert "уже XLSX" in capsys.readouterr().out
    assert main(["convert", str(src), "--to", "mp3", "-o", str(out)]) == 2
