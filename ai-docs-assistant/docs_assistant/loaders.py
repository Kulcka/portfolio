"""Загрузка документов: PDF, DOCX, TXT, MD.

Каждый документ превращается в список разделов :class:`Section` с метаданными
для ссылок: у PDF — номер страницы, у DOCX/MD/TXT — ближайший заголовок.
В DOCX страниц как таковых нет (их расставляет Word при вёрстке), поэтому
ссылка там идёт на раздел.
"""

from __future__ import annotations

import re
from collections.abc import Callable
from dataclasses import dataclass
from pathlib import Path

SUPPORTED_EXTENSIONS = (".pdf", ".docx", ".txt", ".md")

# Меняется при любом изменении правил чтения: индекс тогда пересобирается целиком,
# иначе неизменённые файлы остались бы в базе в старом виде.
LOADER_VERSION = "3"


class DocumentLoadError(Exception):
    """Файл не удалось прочитать; текст ошибки понятен пользователю."""


@dataclass(frozen=True)
class Section:
    text: str
    page: int | None = None
    heading: str | None = None


@dataclass(frozen=True)
class LoadedDocument:
    title: str
    sections: tuple[Section, ...]


def load_document(path: Path) -> LoadedDocument:
    """Прочитать документ. Бросает :class:`DocumentLoadError` с понятной причиной."""
    loader = _LOADERS.get(path.suffix.lower())
    if loader is None:
        raise DocumentLoadError(f"формат {path.suffix or 'без расширения'} не поддерживается")
    try:
        document = loader(path)
    except DocumentLoadError:
        raise
    except Exception as exc:  # noqa: BLE001 - любая ошибка парсера превращается в понятную причину
        raise DocumentLoadError(f"не удалось прочитать файл ({type(exc).__name__}: {exc})") from exc
    if not any(section.text.strip() for section in document.sections):
        hint = " — похоже на скан без текстового слоя, нужен OCR" if path.suffix.lower() == ".pdf" else ""
        raise DocumentLoadError(f"в файле нет текста{hint}")
    return document


# --- PDF ---------------------------------------------------------------------


def _load_pdf(path: Path) -> LoadedDocument:
    from pypdf import PdfReader
    from pypdf.errors import PdfReadError

    try:
        reader = PdfReader(str(path))
    except PdfReadError as exc:
        raise DocumentLoadError(f"повреждённый PDF ({exc})") from exc
    if reader.is_encrypted:
        try:
            decrypted = reader.decrypt("")
        except Exception:  # noqa: BLE001
            decrypted = 0
        if not decrypted:
            raise DocumentLoadError("PDF защищён паролем")

    pages = [clean_text(page.extract_text() or "") for page in reader.pages]
    pages = [unwrap_lines(text) for text in strip_running_lines(pages)]
    sections = [Section(text=text, page=number) for number, text in enumerate(pages, start=1) if text]

    title = ""
    try:
        if reader.metadata and reader.metadata.title:
            title = str(reader.metadata.title).strip()
    except Exception:  # noqa: BLE001 - метаданные необязательны
        title = ""
    return LoadedDocument(title=title or _first_line(sections) or path.stem, sections=tuple(sections))


_WRAP_RE = re.compile(r"(?<![.!?:;…])\n(?=[a-zа-яё0-9(«\"])")


def unwrap_lines(text: str) -> str:
    """Склеить переносы вёрстки PDF: строка без знака конца, а следующая — со строчной буквы."""
    return _WRAP_RE.sub(" ", text)


def strip_running_lines(pages: list[str]) -> list[str]:
    """Убрать колонтитулы: строки, повторяющиеся на большинстве страниц.

    Цифры при сравнении не учитываются, поэтому «Страница 1 из 3» и «Страница 2 из 3»
    считаются одной строкой. Колонтитулом считается только короткая строка (до 120
    символов), которая есть на всех страницах (при 2–4 страницах) или на 60% страниц.
    """
    if len(pages) < 2:
        return pages

    def key(line: str) -> str:
        return re.sub(r"\d+", "#", line.strip().lower())

    counts: dict[str, int] = {}
    for text in pages:
        for line_key in {key(line) for line in text.splitlines() if 0 < len(line.strip()) <= 120}:
            counts[line_key] = counts.get(line_key, 0) + 1
    threshold = len(pages) if len(pages) <= 4 else int(len(pages) * 0.6 + 0.999)
    running = {line_key for line_key, count in counts.items() if count >= threshold}
    if not running:
        return pages
    return [
        clean_text("\n".join(line for line in text.splitlines() if key(line) not in running))
        for text in pages
    ]


# --- DOCX --------------------------------------------------------------------

_HEADING_STYLE_RE = re.compile(r"^(heading|заголовок)\s*(\d+)?$", re.IGNORECASE)


def _load_docx(path: Path) -> LoadedDocument:
    import docx
    from docx.table import Table

    document = docx.Document(str(path))
    title = (document.core_properties.title or "").strip()

    builder = _SectionBuilder()
    for block in document.iter_inner_content():
        if isinstance(block, Table):
            for row in block.rows:
                cells: list[str] = []
                for cell in row.cells:
                    value = " ".join(cell.text.split())
                    # Объединённые ячейки python-docx возвращает повторно — убираем дубли.
                    if value and (not cells or cells[-1] != value):
                        cells.append(value)
                if cells:
                    builder.add_line(" | ".join(cells))
            continue

        text = " ".join(block.text.split())
        if not text:
            continue
        style_name = (block.style.name if block.style is not None else "") or ""
        if style_name.lower() == "title":
            title = title or text
            continue
        if _HEADING_STYLE_RE.match(style_name.strip()):
            builder.start(text)
        else:
            builder.add_line(text)

    sections = builder.finish()
    return LoadedDocument(title=title or _first_line(sections) or path.stem, sections=tuple(sections))


# --- Markdown ----------------------------------------------------------------

_MD_HEADING_RE = re.compile(r"^(#{1,6})\s+(.+?)\s*#*\s*$")
_HTML_COMMENT_RE = re.compile(r"<!--.*?-->", re.DOTALL)
_FRONT_MATTER_RE = re.compile(r"\A---\s*\n.*?\n---\s*\n", re.DOTALL)


def _load_markdown(path: Path) -> LoadedDocument:
    text = read_text_file(path)
    text = _FRONT_MATTER_RE.sub("", text)
    # Комментарии не видны читателю документа — и не должны доходить до модели.
    text = _HTML_COMMENT_RE.sub("", text)

    builder = _SectionBuilder(unwrap=True)
    title = ""
    in_code = False
    for line in text.splitlines():
        if line.strip().startswith("```"):
            in_code = not in_code
            continue
        match = None if in_code else _MD_HEADING_RE.match(line)
        if match:
            heading = _strip_md_inline(match.group(2))
            if len(match.group(1)) == 1 and not title:
                title = heading
                continue
            builder.start(heading)
        else:
            builder.add_line(line.rstrip())
    sections = builder.finish()
    return LoadedDocument(title=title or _first_line(sections) or path.stem, sections=tuple(sections))


def _strip_md_inline(text: str) -> str:
    return re.sub(r"[*_`]+", "", text).strip()


# --- TXT ---------------------------------------------------------------------


def _load_txt(path: Path) -> LoadedDocument:
    text = read_text_file(path)
    builder = _SectionBuilder(unwrap=True)
    lines = text.splitlines()
    title = ""
    for index, line in enumerate(lines):
        stripped = line.strip()
        if not title and stripped:
            title = stripped
            continue
        if _looks_like_txt_heading(stripped, lines[index + 1] if index + 1 < len(lines) else ""):
            builder.start(stripped.rstrip(":"))
        else:
            builder.add_line(line.rstrip())
    sections = builder.finish()
    return LoadedDocument(title=title or path.stem, sections=tuple(sections))


def _looks_like_txt_heading(line: str, next_line: str) -> bool:
    """Заголовок в простом тексте: строка ПРОПИСНЫМИ буквами перед непустым текстом."""
    if not line or len(line) > 100 or not next_line.strip():
        return False
    letters = [ch for ch in line if ch.isalpha()]
    return len(letters) >= 3 and all(ch.isupper() for ch in letters)


def read_text_file(path: Path) -> str:
    """Текст с определением кодировки: UTF-8 (с BOM и без), UTF-16 по BOM, иначе CP1251."""
    raw = path.read_bytes()
    if raw.startswith((b"\xff\xfe", b"\xfe\xff")):
        return raw.decode("utf-16")
    try:
        return raw.decode("utf-8-sig")
    except UnicodeDecodeError:
        return raw.decode("cp1251", errors="replace")


# --- Общее -------------------------------------------------------------------


class _SectionBuilder:
    """Собирает строки в разделы по заголовкам."""

    def __init__(self, *, unwrap: bool = False) -> None:
        self._unwrap = unwrap  # склеивать мягкие переносы строк (Markdown, простой текст)
        self._sections: list[Section] = []
        self._heading: str | None = None
        self._lines: list[str] = []

    def start(self, heading: str) -> None:
        self._flush()
        self._heading = heading

    def add_line(self, line: str) -> None:
        self._lines.append(line)

    def finish(self) -> list[Section]:
        self._flush()
        return self._sections

    def _flush(self) -> None:
        text = clean_text("\n".join(self._lines))
        if self._unwrap:
            text = unwrap_lines(text)
        if text:
            self._sections.append(Section(text=text, heading=self._heading))
        self._lines = []


def clean_text(text: str) -> str:
    """Убрать переносы по слогам, лишние пробелы и пустые строки."""
    text = text.replace("\r\n", "\n").replace("\r", "\n").replace("­", "")
    # «пере-\nсылка» → «пересылка»; «Санкт-\nПетербург» не трогаем (дальше заглавная).
    text = re.sub(r"(\w)-\n([a-zа-яё])", r"\1\2", text)
    lines = [" ".join(line.split()) for line in text.split("\n")]
    text = "\n".join(lines)
    text = re.sub(r"\n{3,}", "\n\n", text)
    return text.strip()


def _first_line(sections: list[Section]) -> str:
    for section in sections:
        for line in section.text.splitlines():
            if line.strip():
                return line.strip()[:120]
    return ""


_LOADERS: dict[str, Callable[[Path], LoadedDocument]] = {
    ".pdf": _load_pdf,
    ".docx": _load_docx,
    ".md": _load_markdown,
    ".txt": _load_txt,
}
