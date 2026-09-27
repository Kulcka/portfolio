"""Выгрузка в CSV, XLSX, JSON и Google Sheets (через подменённый клиент)."""

from __future__ import annotations

import csv
import json
import re
import zipfile
from pathlib import Path
from typing import Any

import gspread
import pytest
from openpyxl import load_workbook

from site_parser.config import Column, GoogleSheetsSpec
from site_parser.exporters import ExportOptions, Table, export_csv, export_json, export_xlsx, protect_csv_formula
from site_parser.gsheets import SheetsError, export_google_sheets

COLUMNS = [
    Column("title", "Название"),
    Column("price", "Цена", "price"),
    Column("in_stock", "В наличии", "availability"),
    Column("qty", "Остаток", "int"),
    Column("tags", "Теги", "list"),
    Column("url", "Ссылка", "url"),
]
ITEMS: list[dict[str, Any]] = [
    {
        "title": "Книга «Первая»; том 1",
        "price": 51.77,
        "in_stock": True,
        "qty": 22,
        "tags": ["a", "b"],
        "url": "https://books.toscrape.com/catalogue/a/index.html",
    },
    {
        "title": '=HYPERLINK("http://evil")',
        "price": None,
        "in_stock": False,
        "qty": 0,
        "tags": [],
        "url": "https://books.toscrape.com/catalogue/b/index.html",
    },
    {
        "title": "Строка с\x0bуправляющим символом",
        "price": 1299.5,
        "in_stock": None,
        "qty": None,
        "tags": None,
        "url": None,
    },
]
OPTS = ExportOptions()


def test_csv_utf8_bom_delimiter_and_values(tmp_path: Path) -> None:
    path = export_csv(tmp_path / "out" / "books.csv", COLUMNS, ITEMS, OPTS)
    raw = path.read_bytes()
    assert raw.startswith(b"\xef\xbb\xbf")  # BOM — Excel сам поймёт UTF-8
    with path.open(encoding="utf-8-sig", newline="") as fh:
        rows = list(csv.reader(fh, delimiter=";"))
    assert rows[0] == ["Название", "Цена", "В наличии", "Остаток", "Теги", "Ссылка"]
    assert rows[1] == ["Книга «Первая»; том 1", "51.77", "да", "22", "a, b", ITEMS[0]["url"]]
    assert rows[2][0] == '\'=HYPERLINK("http://evil")'  # формула экранирована
    assert rows[2][1:4] == ["", "нет", "0"]
    assert len(rows) == 4


def test_csv_custom_delimiter_and_bool_words(tmp_path: Path) -> None:
    opts = ExportOptions(delimiter=",", bool_values=("yes", "no"), list_separator="|")
    path = export_csv(tmp_path / "b.csv", COLUMNS, ITEMS[:1], opts)
    with path.open(encoding="utf-8-sig", newline="") as fh:
        rows = list(csv.reader(fh))
    assert rows[1][2] == "yes" and rows[1][4] == "a|b"


@pytest.mark.parametrize(
    ("value", "expected"),
    [
        ("=1+1", "'=1+1"),
        ("@SUM(A1)", "'@SUM(A1)"),
        ("-5", "-5"),
        ("+7 999 123-45-67", "'+7 999 123-45-67"),
        ("текст", "текст"),
    ],
)
def test_protect_csv_formula(value: str, expected: str) -> None:
    assert protect_csv_formula(value) == expected


def test_xlsx_structure_and_formatting(tmp_path: Path) -> None:
    changes = Table(["Тип", "Ключ", "Было", "Стало"], [["Цена", "A", 10.0, 9.0], ["Новый", "B", None, 5.0]])
    summary = [("Источник", "Books to Scrape"), ("Записей", 3)]
    path = export_xlsx(
        tmp_path / "books.xlsx", COLUMNS, ITEMS, OPTS, sheet_name="Товары", changes=changes, summary=summary
    )

    wb = load_workbook(path)
    assert wb.sheetnames == ["Товары", "Изменения", "Сводка"]
    ws = wb["Товары"]
    assert [c.value for c in ws[1]] == ["Название", "Цена", "В наличии", "Остаток", "Теги", "Ссылка"]
    assert ws.auto_filter.ref == "A1:F4"
    assert ws.freeze_panes == "A2"
    assert ws["A1"].font.bold
    assert ws["B2"].value == 51.77 and ws["B2"].number_format == "#,##0.00"
    assert ws["C2"].value == "да" and ws["C3"].value == "нет"
    assert ws["E2"].value == "a, b"
    assert ws["F2"].hyperlink is not None and ws["F2"].hyperlink.target == ITEMS[0]["url"]
    # строка, похожая на формулу, сохранена как текст
    assert ws["A3"].value == '=HYPERLINK("http://evil")' and ws["A3"].data_type == "s"
    # недопустимый для Excel управляющий символ удалён, а не уронил выгрузку
    assert ws["A4"].value == "Строка суправляющим символом"
    # ширина колонок подобрана по содержимому, но не больше 60
    widths = {col: ws.column_dimensions[col].width for col in "ABCDEF"}
    assert widths["A"] > widths["D"] and max(widths.values()) <= 60

    ch = wb["Изменения"]
    assert [c.value for c in ch[1]] == ["Тип", "Ключ", "Было", "Стало"]
    assert ch["A2"].fill.fgColor.rgb.endswith("FFF2CC")  # цена — жёлтым
    assert wb["Сводка"]["B3"].value == 3


def test_xlsx_empty_changes_sheet(tmp_path: Path) -> None:
    path = export_xlsx(tmp_path / "x.xlsx", COLUMNS, ITEMS[:1], OPTS, changes=Table(["Тип"], []))
    assert load_workbook(path)["Изменения"]["A2"].value == "Изменений нет"


def _full_workbook(tmp_path: Path) -> Path:
    changes = Table(["Тип", "Ключ", "Было", "Стало"], [["Цена", "A", 10.0, 9.0], ["Новый", "B", None, 5.0]])
    summary = [("Источник", "Books to Scrape"), ("Записей", 3)]
    return export_xlsx(
        tmp_path / "books.xlsx", COLUMNS, ITEMS, OPTS, sheet_name="Товары", changes=changes, summary=summary
    )


def test_xlsx_sheet_views_have_no_orphan_pane_selection(tmp_path: Path) -> None:
    """Выделение, привязанное к панели, без самой панели Excel считает повреждением файла.

    Так было на листе «Сводка»: закрепление снимали после установки, openpyxl
    оставлял <selection pane="bottomLeft">, и Excel открывал книгу только
    в режиме восстановления, хотя openpyxl читал её без ошибок.
    """
    path = _full_workbook(tmp_path)
    with zipfile.ZipFile(path) as zf:
        sheets = [n for n in zf.namelist() if re.fullmatch(r"xl/worksheets/sheet\d+\.xml", n)]
        assert len(sheets) == 3
        for name in sheets:
            xml = zf.read(name).decode("utf-8")
            if "<pane " not in xml:
                assert "pane=" not in xml, f"{name}: выделение ссылается на панель, которой нет"
    assert load_workbook(path)["Сводка"].freeze_panes is None


@pytest.mark.excel
def test_xlsx_opens_in_real_excel(tmp_path: Path) -> None:
    """Файл открывается настоящим Excel в обычном режиме, без восстановления."""
    client = pytest.importorskip("win32com.client")
    path = _full_workbook(tmp_path)
    try:
        app = client.DispatchEx("Excel.Application")
    except Exception as exc:  # нет Excel на машине
        pytest.skip(f"Excel недоступен: {exc}")
    app.Visible = False
    app.DisplayAlerts = False
    try:
        # CorruptLoad=0 — обычное открытие: повреждённый файл даёт исключение
        wb = app.Workbooks.Open(str(path.resolve()), UpdateLinks=0, ReadOnly=True, CorruptLoad=0)
        try:
            assert [ws.Name for ws in wb.Worksheets] == ["Товары", "Изменения", "Сводка"]
            data = wb.Worksheets("Товары")
            assert data.Range("B2").Value == 51.77
            assert data.Range("A3").Value == '=HYPERLINK("http://evil")'  # текст, не формула
            assert not data.Range("A3").HasFormula
        finally:
            wb.Close(SaveChanges=False)
    finally:
        app.Quit()


def test_locked_file_is_saved_next_to_it(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    import site_parser.exporters as module

    real_replace = module.os.replace
    target = tmp_path / "books.csv"

    def fake_replace(src: Any, dst: Any) -> None:
        if Path(dst) == target:
            raise PermissionError("файл открыт в Excel")
        real_replace(src, dst)

    monkeypatch.setattr(module.os, "replace", fake_replace)
    saved = export_csv(target, COLUMNS, ITEMS[:1], OPTS)
    assert saved != target and saved.name.startswith("books_") and saved.exists()
    assert not list(tmp_path.glob(".*.tmp*"))


def test_json_keeps_raw_types(tmp_path: Path) -> None:
    path = export_json(tmp_path / "books.json", COLUMNS, ITEMS[:2])
    data = json.loads(path.read_text(encoding="utf-8"))
    assert data[0]["in_stock"] is True and data[0]["tags"] == ["a", "b"] and data[0]["price"] == 51.77
    assert list(data[0]) == [c.name for c in COLUMNS]


# --------------------------------------------------------- Google Sheets


class FakeWorksheet:
    def __init__(self, title: str) -> None:
        self.title = title
        self.calls: list[tuple[str, Any]] = []

    def clear(self) -> None:
        self.calls.append(("clear", None))

    def resize(self, rows: int, cols: int) -> None:
        self.calls.append(("resize", (rows, cols)))

    def update(self, values: list[list[Any]], range_name: str, value_input_option: str) -> None:
        self.calls.append(("update", (values, range_name, value_input_option)))

    def freeze(self, rows: int) -> None:
        self.calls.append(("freeze", rows))

    def format(self, rng: str, fmt: dict[str, Any]) -> None:
        self.calls.append(("format", rng))

    def set_basic_filter(self) -> None:
        self.calls.append(("filter", None))


class FakeSpreadsheet:
    url = "https://docs.google.com/spreadsheets/d/KEY"

    def __init__(self) -> None:
        self.sheets: dict[str, FakeWorksheet] = {"Товары": FakeWorksheet("Товары")}

    def worksheet(self, title: str) -> FakeWorksheet:
        if title not in self.sheets:
            raise gspread.WorksheetNotFound(title)
        return self.sheets[title]

    def add_worksheet(self, title: str, rows: int, cols: int) -> FakeWorksheet:
        self.sheets[title] = FakeWorksheet(title)
        return self.sheets[title]


class FakeClient:
    def __init__(self, spreadsheet: FakeSpreadsheet | None) -> None:
        self.spreadsheet = spreadsheet
        self.opened: list[str] = []

    def open_by_key(self, key: str) -> FakeSpreadsheet:
        self.opened.append(key)
        if self.spreadsheet is None:
            raise gspread.SpreadsheetNotFound("nope")
        return self.spreadsheet

    def open_by_url(self, url: str) -> FakeSpreadsheet:
        return self.open_by_key(url)


@pytest.fixture
def key_file(tmp_path: Path) -> Path:
    path = tmp_path / "sa.json"
    path.write_text(json.dumps({"client_email": "parser@demo.iam.gserviceaccount.com"}), encoding="utf-8")
    return path


def test_google_sheets_writes_raw_values(key_file: Path) -> None:
    spreadsheet = FakeSpreadsheet()
    client = FakeClient(spreadsheet)
    spec = GoogleSheetsSpec(spreadsheet="KEY", worksheet="Товары", changes_worksheet="Изменения")
    url = export_google_sheets(
        spec,
        key_file,
        ["Название", "Цена"],
        [["Книга", 51.77], ["Без цены", None]],
        changes=Table(["Тип"], []),
        client_factory=lambda path: client,
    )
    assert url == FakeSpreadsheet.url and client.opened == ["KEY"]
    data = spreadsheet.sheets["Товары"].calls
    assert data[0] == ("clear", None)
    assert data[1] == ("resize", (3, 2))
    values, rng, mode = data[2][1]
    assert values == [["Название", "Цена"], ["Книга", 51.77], ["Без цены", ""]]
    assert (rng, mode) == ("A1", "RAW")  # RAW: строки с сайта не исполняются как формулы
    changes_sheet = spreadsheet.sheets["Изменения"]  # лист создан, раз его не было
    assert changes_sheet.calls[2][1][0] == [["Тип"], ["Изменений нет"]]


def test_google_sheets_without_key_or_access(key_file: Path, tmp_path: Path) -> None:
    spec = GoogleSheetsSpec(spreadsheet="KEY")
    with pytest.raises(SheetsError, match="GOOGLE_SERVICE_ACCOUNT_FILE"):
        export_google_sheets(spec, None, ["a"], [])
    with pytest.raises(SheetsError, match="не найден"):
        export_google_sheets(spec, tmp_path / "missing.json", ["a"], [])
    with pytest.raises(SheetsError, match=re.escape("parser@demo.iam.gserviceaccount.com")):
        export_google_sheets(spec, key_file, ["a"], [], client_factory=lambda path: FakeClient(None))
