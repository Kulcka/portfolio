"""Прогон целиком без сети: сбор → картинки → история → CSV/XLSX/JSON → отчёт → уведомление."""

from __future__ import annotations

import dataclasses
import json
from pathlib import Path
from typing import Any, ClassVar

import pytest
from openpyxl import load_workbook

from site_parser.cli import main
from site_parser.config import GoogleSheetsSpec, SiteConfig
from site_parser.fetcher import Fetcher
from site_parser.pipeline import EXIT_OK, PipelineError, RunOptions, run_pipeline
from site_parser.settings import EnvSettings
from site_parser.storage import SnapshotStore
from tests.conftest import (
    BOOKS,
    CONFIGS,
    FakeClock,
    FakeSession,
    books_product_page,
    books_session,
    fixture_bytes,
    html_response,
)
from tests.test_exporters import FakeClient, FakeSpreadsheet

TOKEN = "987654321:AAF-another-fake-token-for-tests-000"


class CapturingNotifier:
    sent: ClassVar[list[str]] = []

    def __init__(self, token: str, chat_id: str) -> None:
        assert token == TOKEN and chat_id == "42"

    def send(self, text: str) -> None:
        CapturingNotifier.sent.append(text)


@pytest.fixture
def workdir(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> Path:
    monkeypatch.chdir(tmp_path)  # относительные пути конфига (output/, state/) — во временной папке
    CapturingNotifier.sent = []
    return tmp_path


def _fetcher(config: SiteConfig, session: FakeSession) -> Fetcher:
    clock = FakeClock()
    return Fetcher(config.request, session=session, sleep=clock.sleep, clock=clock.time)


def _changed_site() -> FakeSession:
    """Тот же каталог, но: у одной книги новая цена, одна книга «переехала» на новый адрес,
    последняя книга на второй странице снята с продажи, у одной книги изменился остаток."""
    session = books_session()
    page1 = fixture_bytes("books/catalogue_page-1.html")
    page1 = page1.replace("£53.74".encode(), "£49.99".encode())  # Tipping the Velvet
    page1 = page1.replace(b"soumission_998/index.html", b"soumission-2nd-edition_1001/index.html")
    session.routes[f"{BOOKS}/catalogue/page-1.html"] = html_response(page1)
    page2 = fixture_bytes("books/catalogue_page-50.html").decode("utf-8")
    last = page2.rindex('<li class="col-xs-6 col-sm-4 col-md-3 col-lg-3">')
    end = page2.index("</li>", last) + len("</li>")
    session.routes[f"{BOOKS}/catalogue/page-2.html"] = html_response(page2[:last] + page2[end:])
    sharp_url = f"{BOOKS}/catalogue/sharp-objects_997/index.html"

    def sharp_objects(url: str) -> Any:
        response = books_product_page(url)
        response.content = response.content.replace(b"22 available", b"3 available")
        return response

    session.routes[sharp_url] = sharp_objects
    return session


def test_two_runs_produce_files_and_change_report(books_config: SiteConfig, workdir: Path) -> None:
    env = EnvSettings(telegram_token=TOKEN, telegram_chat_id="42")
    options = RunOptions(max_images=3, notify_preview=workdir / "preview.txt")

    first = run_pipeline(
        books_config, env, options, fetcher=_fetcher(books_config, books_session()), notifier_factory=CapturingNotifier
    )  # type: ignore[arg-type]
    assert first.exit_code == EXIT_OK
    assert len(first.scrape.items) == 40 and first.diff.is_first_run
    out = workdir / "output" / "books_toscrape"
    assert {p.name for p in first.files} == {"books.csv", "books.xlsx", "books.json", "changes.md"}
    assert len(list((out / "images").glob("*.jpg"))) == 3
    assert "Первый прогон" in (out / "changes.md").read_text(encoding="utf-8")
    assert CapturingNotifier.sent == []  # only_changes: первый прогон — не изменение
    assert "Первый прогон: собрано записей 40." in (workdir / "preview.txt").read_text(encoding="utf-8")

    second = run_pipeline(
        books_config, env, options, fetcher=_fetcher(books_config, _changed_site()), notifier_factory=CapturingNotifier
    )  # type: ignore[arg-type]
    diff = second.diff
    assert [c.title for c in diff.price_changes] == ["Tipping the Velvet"]
    assert (diff.price_changes[0].old, diff.price_changes[0].new) == (53.74, 49.99)
    assert [i["url"] for i in diff.new] == [f"{BOOKS}/catalogue/soumission-2nd-edition_1001/index.html"]
    assert {i["url"].split("/")[-2] for i in diff.gone} == {"soumission_998", "1000-places-to-see-before-you-die_1"}
    assert [(c.field, c.old, c.new) for c in diff.field_changes] == [("stock_qty", 22, 3)]

    report = (out / "changes.md").read_text(encoding="utf-8")
    assert "## Изменилась цена (1)" in report and "53.74 | 49.99 | −7.0%" in report
    assert "## Новые (1)" in report and "## Пропали (2)" in report
    assert len(list((out / "reports").glob("changes_*.md"))) >= 1

    assert len(CapturingNotifier.sent) == 1
    assert "Новые: 1 · Цена: 1 · Пропали: 2 · Другое: 1" in CapturingNotifier.sent[0]

    wb = load_workbook(out / "books.xlsx")
    assert wb.sheetnames == ["Товары", "Изменения", "Сводка"]
    kinds = [row[0].value for row in wb["Изменения"].iter_rows(min_row=2)]
    assert sorted(kinds) == sorted(["Новый", "Цена", "Изменение", "Пропал", "Пропал"])
    header = [c.value for c in wb["Товары"][1]]
    assert header[:5] == ["Название", "Цена", "Валюта", "В наличии", "Рейтинг"]
    assert "Характеристика: Product Type" in header and header[-2:] == ["Ссылка", "Файл картинки"]

    data = json.loads((out / "books.json").read_text(encoding="utf-8"))
    assert len(data) == 39 and data[0]["upc"] and data[0]["image_file"].startswith("images/")

    runs = SnapshotStore(books_config.storage.path).list_runs("books_toscrape")
    assert [r.items_count for r in runs] == [39, 40]


def test_nothing_scraped_keeps_previous_files(books_config: SiteConfig, workdir: Path) -> None:
    run_pipeline(
        books_config,
        EnvSettings(),
        RunOptions(max_pages=1, max_images=0),
        fetcher=_fetcher(books_config, books_session()),
    )
    csv_path = workdir / "output" / "books_toscrape" / "books.csv"
    before = csv_path.read_bytes()
    broken = FakeSession(
        {
            f"{BOOKS}/robots.txt": html_response("", status=404),
            f"{BOOKS}/catalogue/page-1.html": html_response("<html>Ведутся работы</html>"),
        }
    )
    with pytest.raises(PipelineError, match="не собрано ни одной записи"):
        run_pipeline(books_config, EnvSettings(), RunOptions(), fetcher=_fetcher(books_config, broken))
    assert csv_path.read_bytes() == before
    assert len(SnapshotStore(books_config.storage.path).list_runs("books_toscrape")) == 1


def test_google_sheets_export_and_failure_is_a_warning(books_config: SiteConfig, workdir: Path) -> None:
    key = workdir / "sa.json"
    key.write_text(json.dumps({"client_email": "bot@demo.iam.gserviceaccount.com"}), encoding="utf-8")
    config = dataclasses.replace(
        books_config,
        output=dataclasses.replace(
            books_config.output, google_sheets=GoogleSheetsSpec(spreadsheet="KEY", worksheet="Товары")
        ),
    )
    spreadsheet = FakeSpreadsheet()
    outcome = run_pipeline(
        config,
        EnvSettings(google_credentials_file=key),
        RunOptions(max_pages=1, max_images=0),
        fetcher=_fetcher(config, books_session()),
        sheets_client_factory=lambda p: FakeClient(spreadsheet),
    )
    assert outcome.sheet_url == FakeSpreadsheet.url and outcome.exit_code == EXIT_OK
    values = spreadsheet.sheets["Товары"].calls[2][1][0]
    assert values[0][0] == "Название" and len(values) == 21

    failed = run_pipeline(
        config,
        EnvSettings(google_credentials_file=None),
        RunOptions(max_pages=1, max_images=0),
        fetcher=_fetcher(config, books_session()),
    )
    assert failed.exit_code == 3 and "Google Sheets" in failed.warnings[0]
    assert (workdir / "output" / "books_toscrape" / "books.xlsx").exists()  # локальные файлы на месте


def test_cli_run_check_and_history(
    workdir: Path, monkeypatch: pytest.MonkeyPatch, capsys: pytest.CaptureFixture[str]
) -> None:
    def fake_from_settings(settings: Any, *, use_cache: bool | None = None) -> Fetcher:
        clock = FakeClock()
        return Fetcher(settings, session=books_session(), sleep=clock.sleep, clock=clock.time)

    monkeypatch.setattr(Fetcher, "from_settings", staticmethod(fake_from_settings))
    env_file = workdir / "empty.env"
    env_file.write_text("", encoding="utf-8")
    config = str(CONFIGS / "books_toscrape.yaml")

    assert (
        main(["run", config, "--env-file", str(env_file), "--max-pages", "1", "--max-images", "0", "--no-notify"]) == 0
    )
    assert (workdir / "output" / "books_toscrape" / "books.xlsx").exists()
    assert (workdir / "logs" / "site_parser.log").exists()

    assert main(["check", config, "--env-file", str(env_file), "--limit", "2"]) == 0
    printed = capsys.readouterr().out
    assert "--- Запись 2 ---" in printed and "Название: A Light in the Attic" in printed

    assert main(["history", config, "--env-file", str(env_file)]) == 0
    assert "20" in capsys.readouterr().out

    assert main(["run", str(workdir / "missing.yaml"), "--env-file", str(env_file)]) == 2
