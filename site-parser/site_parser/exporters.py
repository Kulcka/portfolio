"""Выгрузка в CSV, XLSX и JSON.

* CSV — UTF-8 с BOM (Excel сразу понимает кириллицу), разделитель из конфига
  (по умолчанию «;» — так CSV открывается двойным щелчком в русском Excel).
* XLSX — лист с данными (заголовки, автофильтр, закреплённая шапка, ширина
  колонок по содержимому, формат цен, кликабельные ссылки), лист
  «Изменения» и лист «Сводка».
* JSON — «сырые» значения для программ.

Файлы пишутся атомарно: сначала во временный, потом замена. Если файл открыт
в Excel (Windows его блокирует), результат сохраняется рядом с отметкой
времени в имени — прогон по расписанию не падает.

Строки, похожие на формулы (``=HYPERLINK(...)`` со страницы сайта), в Excel
записываются как текст, в CSV — с апострофом в начале: данные с чужого сайта
не должны исполняться как формулы.
"""

from __future__ import annotations

import csv
import json
import logging
import os
import re
from collections.abc import Callable, Iterable, Sequence
from dataclasses import dataclass
from datetime import datetime
from pathlib import Path
from typing import Any

from openpyxl import Workbook
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from site_parser.config import Column

log = logging.getLogger(__name__)

EXCEL_MAX_CELL = 32_767
_NUMERIC_RE = re.compile(r"^[-+]?\d[\d\s.,]*$")
_FORMULA_PREFIXES = ("=", "+", "-", "@", "\t", "\r")

HEADER_FONT = Font(bold=True, color="FFFFFF")
HEADER_FILL = PatternFill("solid", fgColor="2F5597")
CHANGE_FILLS = {
    "Новый": PatternFill("solid", fgColor="E2EFDA"),
    "Цена": PatternFill("solid", fgColor="FFF2CC"),
    "Изменение": PatternFill("solid", fgColor="DDEBF7"),
    "Пропал": PatternFill("solid", fgColor="FCE4D6"),
}
NUMBER_FORMATS = {"price": "#,##0.00", "float": "#,##0.00", "int": "0"}


class ExportError(Exception):
    """Не удалось записать файл выгрузки."""


@dataclass(frozen=True)
class ExportOptions:
    delimiter: str = ";"
    list_separator: str = ", "
    bool_values: tuple[str, str] = ("да", "нет")


@dataclass(frozen=True)
class Table:
    """Произвольная таблица (лист «Изменения»)."""

    header: Sequence[str]
    rows: Sequence[Sequence[Any]]


def human_value(value: Any, opts: ExportOptions) -> Any:
    """Значение для таблиц: списки — через разделитель, bool — да/нет."""
    if value is None:
        return None
    if isinstance(value, bool):
        return opts.bool_values[0] if value else opts.bool_values[1]
    if isinstance(value, (list, tuple)):
        parts = [human_value(v, opts) for v in value]
        return opts.list_separator.join(str(p) for p in parts if p not in (None, ""))
    if isinstance(value, (int, float, str)):
        return value
    return str(value)


def protect_csv_formula(text: str) -> str:
    """Экранировать строку, которую Excel принял бы за формулу."""
    if text.startswith(_FORMULA_PREFIXES) and not _NUMERIC_RE.match(text):
        return "'" + text
    return text


def rows_for(items: Iterable[dict[str, Any]], columns: Sequence[Column], opts: ExportOptions) -> list[list[Any]]:
    return [[human_value(item.get(col.name), opts) for col in columns] for item in items]


def write_atomic(path: Path, write: Callable[[Path], None]) -> Path:
    """Записать файл через временный. Если целевой занят — сохранить рядом с датой."""
    path.parent.mkdir(parents=True, exist_ok=True)
    tmp = path.with_name(f".{path.stem}.tmp{path.suffix}")
    try:
        write(tmp)
    except OSError as exc:
        tmp.unlink(missing_ok=True)
        raise ExportError(f"не удалось записать {path}: {exc.strerror or exc}") from None
    try:
        os.replace(tmp, path)
        return path
    except PermissionError:
        alt = path.with_name(f"{path.stem}_{datetime.now():%Y%m%d_%H%M%S}{path.suffix}")
        os.replace(tmp, alt)
        log.warning("Файл %s занят другой программой (открыт в Excel?) — сохранено в %s", path, alt.name)
        return alt


# --------------------------------------------------------------------- CSV


def export_csv(path: Path, columns: Sequence[Column], items: Sequence[dict[str, Any]], opts: ExportOptions) -> Path:
    def write(target: Path) -> None:
        with target.open("w", encoding="utf-8-sig", newline="") as fh:
            writer = csv.writer(fh, delimiter=opts.delimiter, lineterminator="\r\n")
            writer.writerow([col.header for col in columns])
            for row in rows_for(items, columns, opts):
                writer.writerow(["" if v is None else protect_csv_formula(v) if isinstance(v, str) else v for v in row])

    return write_atomic(path, write)


# -------------------------------------------------------------------- JSON


def export_json(path: Path, columns: Sequence[Column], items: Sequence[dict[str, Any]]) -> Path:
    ordered = [{col.name: item.get(col.name) for col in columns} for item in items]

    def write(target: Path) -> None:
        target.write_text(json.dumps(ordered, ensure_ascii=False, indent=2, default=str), encoding="utf-8")

    return write_atomic(path, write)


# -------------------------------------------------------------------- XLSX


def _clean_text(value: str) -> str:
    value = ILLEGAL_CHARACTERS_RE.sub("", value)
    return value if len(value) <= EXCEL_MAX_CELL else value[: EXCEL_MAX_CELL - 1] + "…"


def _set_cell(ws: Worksheet, row: int, col: int, value: Any) -> Any:
    cell = ws.cell(row=row, column=col)
    if isinstance(value, str):
        value = _clean_text(value)
        cell.value = value
        if value.startswith("="):
            cell.data_type = "s"  # текст, а не формула
    else:
        cell.value = value
    return cell


def _write_header(ws: Worksheet, header: Sequence[str], *, freeze: bool = True) -> None:
    for idx, title in enumerate(header, start=1):
        cell = _set_cell(ws, 1, idx, title)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(vertical="center", wrap_text=True)
    # Закреплять только здесь: «freeze_panes = None» после закрепления оставляет
    # <selection pane="bottomLeft"> без панели, и Excel считает файл повреждённым.
    if freeze:
        ws.freeze_panes = "A2"
    ws.row_dimensions[1].height = 30


def _autosize(ws: Worksheet, *, max_width: int = 60, sample_rows: int = 1000) -> None:
    """Ширина колонок по самому длинному значению (по первым ``sample_rows`` строкам)."""
    widths: dict[int, int] = {}
    for row in ws.iter_rows(min_row=1, max_row=min(ws.max_row, sample_rows + 1)):
        for cell in row:
            if cell.value is None:
                continue
            value = cell.value
            text = f"{value:,.2f}" if isinstance(value, float) else str(value)
            longest = max((len(line) for line in text.splitlines()), default=0)
            widths[cell.column] = max(widths.get(cell.column, 0), longest)
    for col_idx, width in widths.items():
        ws.column_dimensions[get_column_letter(col_idx)].width = max(8, min(width + 2, max_width))


def _fill_data_sheet(
    ws: Worksheet, columns: Sequence[Column], items: Sequence[dict[str, Any]], opts: ExportOptions
) -> None:
    _write_header(ws, [col.header for col in columns])
    for r, row in enumerate(rows_for(items, columns, opts), start=2):
        for c, (col, value) in enumerate(zip(columns, row, strict=True), start=1):
            cell = _set_cell(ws, r, c, value)
            if value is None:
                continue
            if col.kind in NUMBER_FORMATS and isinstance(value, (int, float)):
                cell.number_format = NUMBER_FORMATS[col.kind]
            elif col.kind == "url" and isinstance(value, str) and value.startswith(("http://", "https://")):
                cell.hyperlink = value
                cell.style = "Hyperlink"
    ws.auto_filter.ref = f"A1:{get_column_letter(max(len(columns), 1))}{max(len(items) + 1, 1)}"
    _autosize(ws)


def _fill_changes_sheet(ws: Worksheet, table: Table) -> None:
    _write_header(ws, list(table.header))
    for r, row in enumerate(table.rows, start=2):
        fill = CHANGE_FILLS.get(str(row[0])) if row else None
        for c, value in enumerate(row, start=1):
            cell = _set_cell(ws, r, c, value)
            if fill is not None:
                cell.fill = fill
            if isinstance(value, float):
                cell.number_format = "#,##0.00"
            if isinstance(value, str) and value.startswith(("http://", "https://")):
                cell.hyperlink = value
                cell.font = Font(color="0563C1", underline="single")
    if not table.rows:
        _set_cell(ws, 2, 1, "Изменений нет")
    ws.auto_filter.ref = f"A1:{get_column_letter(len(table.header))}{max(len(table.rows) + 1, 1)}"
    _autosize(ws)


def _fill_summary_sheet(ws: Worksheet, summary: Sequence[tuple[str, Any]]) -> None:
    _write_header(ws, ["Показатель", "Значение"], freeze=False)
    for r, (name, value) in enumerate(summary, start=2):
        _set_cell(ws, r, 1, name).font = Font(bold=True)
        _set_cell(ws, r, 2, value).alignment = Alignment(horizontal="left")
    _autosize(ws, max_width=90)


def export_xlsx(
    path: Path,
    columns: Sequence[Column],
    items: Sequence[dict[str, Any]],
    opts: ExportOptions,
    *,
    sheet_name: str = "Данные",
    changes: Table | None = None,
    summary: Sequence[tuple[str, Any]] | None = None,
) -> Path:
    wb = Workbook()
    ws = wb.active
    assert ws is not None
    ws.title = sheet_name[:31]
    _fill_data_sheet(ws, columns, items, opts)
    if changes is not None:
        _fill_changes_sheet(wb.create_sheet("Изменения"), changes)
    if summary:
        _fill_summary_sheet(wb.create_sheet("Сводка"), summary)

    def write(target: Path) -> None:
        wb.save(target)

    return write_atomic(path, write)
