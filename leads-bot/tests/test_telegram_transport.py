"""Транспорт Telegram через подменённую сессию aiogram — без сети."""

from __future__ import annotations

import io
from collections.abc import AsyncGenerator
from datetime import datetime, timezone
from typing import Any

import pytest
from aiogram import Bot
from aiogram.client.session.base import BaseSession
from aiogram.methods import (
    AnswerCallbackQuery,
    EditMessageReplyMarkup,
    GetMe,
    SendDocument,
    SendMessage,
    TelegramMethod,
)
from aiogram.types import (
    CallbackQuery,
    Chat,
    Contact,
    InlineKeyboardMarkup,
    Message,
    ReplyKeyboardMarkup,
    ReplyKeyboardRemove,
    Update,
    User,
)
from openpyxl import load_workbook

from leads_bot.config import BotConfig
from leads_bot.core.service import LeadService, Payload
from leads_bot.core.storage import Storage
from leads_bot.transports.common import AdminAccess
from leads_bot.transports.telegram import TelegramTransport
from tests.helpers import FakeClock

FAKE_TOKEN = "123456:TEST-TOKEN-NOT-REAL"
ADMIN_CHAT = -100500
ADMIN_USER = 7
CLIENT = 1001


class RecordingSession(BaseSession):
    """Сессия aiogram, которая ничего не шлёт в сеть, а записывает запросы."""

    def __init__(self) -> None:
        super().__init__()
        self.requests: list[TelegramMethod[Any]] = []
        self._message_id = 100

    async def make_request(self, bot: Bot, method: TelegramMethod[Any], timeout: int | None = None) -> Any:
        self.requests.append(method)
        if isinstance(method, (SendMessage, SendDocument)):
            self._message_id += 1
            return Message(
                message_id=self._message_id,
                date=datetime.now(timezone.utc),
                chat=Chat(id=int(method.chat_id), type="private"),
                text=getattr(method, "text", None),
            )
        if isinstance(method, GetMe):
            return User(id=123456, is_bot=True, first_name="Demo", username="demo_leads_bot")
        return True

    async def stream_content(self, *args: Any, **kwargs: Any) -> AsyncGenerator[bytes, None]:
        yield b""

    async def close(self) -> None:
        pass

    def sent(self, kind: type) -> list[Any]:
        return [r for r in self.requests if isinstance(r, kind)]

    def clear(self) -> None:
        self.requests.clear()


def _user(user_id: int) -> User:
    return User(id=user_id, is_bot=False, first_name="Иван", last_name="Петров", username="ivan")


class Harness:
    def __init__(self, service: LeadService) -> None:
        self.session = RecordingSession()
        self.transport = TelegramTransport(
            FAKE_TOKEN,
            service,
            AdminAccess(chat_id=ADMIN_CHAT, user_ids=frozenset({ADMIN_USER})),
            session=self.session,
        )
        service.add_notifier(self.transport.notifier())  # type: ignore[arg-type]
        self._update_id = 0

    async def feed(self, **kwargs: Any) -> None:
        self._update_id += 1
        await self.transport.dispatcher.feed_update(self.transport.bot, Update(update_id=self._update_id, **kwargs))

    async def message(
        self, text: str | None = None, *, chat_id: int = CLIENT, user_id: int = CLIENT, chat_type: str = "private",
        contact: Contact | None = None,
    ) -> None:
        await self.feed(
            message=Message(
                message_id=self._update_id + 1,
                date=datetime.now(timezone.utc),
                chat=Chat(id=chat_id, type=chat_type),
                from_user=_user(user_id),
                text=text,
                contact=contact,
            )
        )

    async def press(self, payload: str, *, user_id: int = CLIENT) -> None:
        await self.feed(
            callback_query=CallbackQuery(
                id=f"cb{self._update_id}",
                from_user=_user(user_id),
                chat_instance="ci",
                data=payload,
                message=Message(
                    message_id=50, date=datetime.now(timezone.utc), chat=Chat(id=user_id, type="private"), text="..."
                ),
            )
        )


@pytest.fixture
def harness(config: BotConfig, storage: Storage, clock: FakeClock) -> Harness:
    return Harness(LeadService(config, storage, clock=clock))


async def test_start_shows_consent_with_policy(harness: Harness, config: BotConfig) -> None:
    await harness.message("/start")
    [sent] = harness.session.sent(SendMessage)
    assert sent.chat_id == CLIENT and "согласие" in sent.text
    markup = sent.reply_markup
    assert isinstance(markup, InlineKeyboardMarkup)
    assert markup.inline_keyboard[0][0].url == config.privacy.policy_url
    assert [b.callback_data for b in markup.inline_keyboard[1]] == [Payload.CONSENT_YES, Payload.CONSENT_NO]


async def test_full_flow_notifies_admin_chat(harness: Harness, storage: Storage) -> None:
    await harness.message("/start")
    await harness.press(Payload.CONSENT_YES)
    # Нажатие подтверждено, кнопки под нажатым сообщением убраны.
    assert harness.session.sent(AnswerCallbackQuery) and harness.session.sent(EditMessageReplyMarkup)
    await harness.press(Payload.SERVICE_PREFIX + "electric")
    await harness.message("Иван")

    phone_prompt = harness.session.sent(SendMessage)[-1]
    assert isinstance(phone_prompt.reply_markup, ReplyKeyboardMarkup)
    assert phone_prompt.reply_markup.keyboard[0][0].request_contact is True

    harness.session.clear()
    await harness.message(contact=Contact(phone_number="79991234567", first_name="Иван", user_id=CLIENT))
    saved, ask_comment = harness.session.sent(SendMessage)
    assert isinstance(saved.reply_markup, ReplyKeyboardRemove) and "+7 999 123-45-67" in saved.text
    assert isinstance(ask_comment.reply_markup, InlineKeyboardMarkup)

    await harness.press(Payload.SKIP_COMMENT)
    harness.session.clear()
    await harness.press(Payload.SEND)

    [lead] = storage.list_leads()
    assert (lead.phone, lead.service_id, lead.username) == ("+79991234567", "electric", "ivan")
    assert lead.admin_notified
    admin, thanks = harness.session.sent(SendMessage)
    assert admin.chat_id == ADMIN_CHAT and "Новая заявка №1" in admin.text and "tg://user?id=1001" in admin.text
    assert thanks.chat_id == CLIENT and "№1" in thanks.text


async def test_cancel_and_forget_commands(harness: Harness, storage: Storage, config: BotConfig) -> None:
    await harness.message("/start")
    await harness.message("/cancel")
    assert harness.session.sent(SendMessage)[-1].text == config.texts.cancelled
    await harness.message("/forget")
    assert "заявок: 0" in harness.session.sent(SendMessage)[-1].text


async def test_stats_only_for_admins(harness: Harness) -> None:
    await harness.message("/stats", chat_id=ADMIN_CHAT, user_id=555, chat_type="supergroup")
    assert "Всего: <b>0</b>" in harness.session.sent(SendMessage)[-1].text

    harness.session.clear()
    await harness.message("/stats", chat_id=ADMIN_USER, user_id=ADMIN_USER)  # админ в личке
    assert "Всего" in harness.session.sent(SendMessage)[-1].text

    harness.session.clear()
    await harness.message("/stats")  # обычный клиент — статистику не видит
    [reply] = harness.session.sent(SendMessage)
    assert "Всего" not in reply.text and "/start" in reply.text


async def test_export_sends_excel(harness: Harness, config: BotConfig) -> None:
    await harness.message("/export", chat_id=ADMIN_CHAT, user_id=555, chat_type="group")
    assert harness.session.sent(SendMessage)[-1].text == "Заявок пока нет."

    await harness.message("/start")
    await harness.press(Payload.CONSENT_YES)
    await harness.press(Payload.SERVICE_PREFIX + "design")
    await harness.message("Иван")
    await harness.message("+7 999 123-45-67")
    await harness.message("Нужен проект")
    await harness.press(Payload.SEND)

    await harness.message("/export", chat_id=ADMIN_CHAT, user_id=555, chat_type="group")
    [document] = harness.session.sent(SendDocument)
    assert document.chat_id == ADMIN_CHAT and document.caption == "Выгрузка заявок: 1 шт."
    assert document.document.filename.endswith(".xlsx")
    ws = load_workbook(io.BytesIO(document.document.data))["Заявки"]
    assert ws.cell(row=2, column=7).value == "Нужен проект"


async def test_group_messages_ignored_and_id_command(harness: Harness) -> None:
    await harness.message("Хочу ремонт", chat_id=-42, user_id=CLIENT, chat_type="group")
    assert harness.session.sent(SendMessage) == []
    await harness.message("/id", chat_id=-42, user_id=CLIENT, chat_type="group")
    assert "-42" in harness.session.sent(SendMessage)[-1].text


async def test_setup_commands(harness: Harness) -> None:
    await harness.transport.setup_commands()
    assert len([r for r in harness.session.requests if type(r).__name__ == "SetMyCommands"]) == 2
