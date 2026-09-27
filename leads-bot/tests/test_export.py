"""Выгрузка в Excel."""

from __future__ import annotations

import io
from datetime import datetime

from openpyxl import load_workbook

from leads_bot.core.export import COLUMNS
from leads_bot.core.models import Channel, UserRef
from leads_bot.core.service import LeadService
from leads_bot.core.storage import Storage
from tests.helpers import FakeClock, fill_lead


async def test_export_contents(service: LeadService, user: UserRef, clock: FakeClock) -> None:
    await fill_lead(service, user, service_id="electric", name="Иван", phone="8 999 123 45 67", comment="Нужна смета")
    clock.advance(minutes=5)
    max_user = UserRef(channel=Channel.MAX, user_id=77, display_name="Мария")
    await fill_lead(service, max_user, service_id="design", name="Мария", phone="+7 916 000-11-22", comment=None)

    export = service.export()
    assert export.count == 2
    assert export.filename == "leads_2026-09-27_1205.xlsx"  # время по Москве

    wb = load_workbook(io.BytesIO(export.content))
    assert wb.sheetnames == ["Заявки", "Сводка"]
    ws = wb["Заявки"]
    assert [c.value for c in ws[1]] == [title for title, _ in COLUMNS]
    assert ws.freeze_panes == "A2"

    first = {title: cell for (title, _), cell in zip(COLUMNS, ws[2], strict=True)}
    assert first["№"].value == 1
    assert first["Дата и время"].value == datetime(2026, 9, 27, 12, 0)  # местное время без часового пояса
    assert first["Дата и время"].number_format == "DD.MM.YYYY HH:MM"
    assert first["Канал"].value == "Telegram"
    assert first["Услуга"].value == "Электрика"
    assert first["Телефон"].value == "+79991234567" and first["Телефон"].data_type == "s"
    assert first["Комментарий"].value == "Нужна смета"
    assert first["Профиль в мессенджере"].value == "@ivan"
    assert first["ID пользователя"].value == "1001"
    assert first["Политика (версия)"].value == "https://example.com/privacy (2026-09-01)"
    assert first["Уведомление отправлено"].value == "да"

    second = [c.value for c in ws[3]]
    assert second[2:7] == ["MAX", "Дизайн-проект", "Мария", "+79160001122", None]  # пустой комментарий
    assert second[7] == "Мария"  # без username — имя профиля

    summary = wb["Сводка"]
    rows = [tuple(c.value for c in row) for row in summary.iter_rows()]
    assert ("Всего", 2) in rows
    assert ("Электрика", 1) in rows and ("Дизайн-проект", 1) in rows


async def test_export_neutralizes_formulas_and_control_chars(service: LeadService, user: UserRef) -> None:
    await fill_lead(service, user, comment='=HYPERLINK("http://evil.example","жми")\x07')
    ws = load_workbook(io.BytesIO(service.export().content))["Заявки"]
    comment = ws.cell(row=2, column=7)
    assert comment.data_type == "s"  # строка, а не формула
    assert comment.value == '=HYPERLINK("http://evil.example","жми")'  # управляющий символ удалён


async def test_export_empty(service: LeadService, storage: Storage) -> None:
    export = service.export()
    assert export.count == 0
    ws = load_workbook(io.BytesIO(export.content))["Заявки"]
    assert ws.max_row == 1
