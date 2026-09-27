"""Транспорт MAX через подменённый API-клиент — без сети.

Формат событий взят из документации MAX (dev.max.ru/docs-api) и официального
Go-клиента (schemes.go): bot_started, message_created, message_callback.
"""

from __future__ import annotations

import asyncio
import io
import json
from typing import Any

import pytest
from openpyxl import load_workbook

from leads_bot.config import BotConfig
from leads_bot.core.models import Channel
from leads_bot.core.service import LeadService, Payload
from leads_bot.core.storage import Storage
from leads_bot.transports.common import ADMIN_EXPORT_FAILED, AdminAccess
from leads_bot.transports.max_api import MaxApiError, MaxAuthError
from leads_bot.transports.max_bot import MARKER_KEY, MaxTransport, contact_phone, phone_from_vcf
from tests.helpers import FakeClock

CLIENT = 2002
DIALOG = 5005
ADMIN_CHAT = -777


class FakeMaxApi:
    def __init__(self) -> None:
        self.sent: list[dict[str, Any]] = []
        self.answers: list[dict[str, Any]] = []
        self.uploads: list[tuple[str, bytes]] = []
        self.commands: list[tuple[str, str]] | None = None
        self.batches: list[Any] = []
        self.polled_markers: list[int | None] = []
        self.subscribed: dict[str, Any] | None = None
        self.fail_upload = False
        self.stop = asyncio.Event()
        self._mid = 0

    async def get_me(self) -> dict[str, Any]:
        return {"user_id": 1, "name": "Demo", "username": "demo_leads_bot", "is_bot": True}

    async def set_commands(self, commands: Any) -> None:
        self.commands = list(commands)

    async def get_updates(self, *, marker: int | None = None, timeout: int = 30, limit: int = 100, types: Any = ()):  # type: ignore[no-untyped-def]
        self.polled_markers.append(marker)
        if not self.batches:
            self.stop.set()
            return [], marker
        item = self.batches.pop(0)
        if isinstance(item, BaseException):
            raise item
        return item

    async def send_message(self, *, text: str, chat_id: int | None = None, user_id: int | None = None,
                           attachments: list[dict[str, Any]] | None = None, text_format: str | None = "html") -> dict[str, Any]:
        self._mid += 1
        self.sent.append({"text": text, "chat_id": chat_id, "user_id": user_id, "attachments": attachments})
        return {"message": {"body": {"mid": f"mid.{self._mid}", "seq": self._mid, "text": text}}}

    async def answer_callback(self, callback_id: str, *, notification: str | None = None,
                              message: dict[str, Any] | None = None) -> None:
        self.answers.append({"callback_id": callback_id, "notification": notification, "message": message})

    async def upload_file(self, filename: str, content: bytes) -> str:
        if self.fail_upload:
            raise MaxApiError(500, None, "upload failed")
        self.uploads.append((filename, content))
        return "file-token-1"

    async def get_subscriptions(self) -> list[dict[str, Any]]:
        return []

    async def subscribe(self, url: str, *, update_types: Any, secret: str | None = None) -> None:
        self.subscribed = {"url": url, "update_types": list(update_types), "secret": secret}

    def last_mid(self) -> str:
        return f"mid.{self._mid}"


def user_obj(user_id: int = CLIENT) -> dict[str, Any]:
    return {"user_id": user_id, "first_name": "Мария", "last_name": "", "name": "Мария", "is_bot": False}


def bot_started() -> dict[str, Any]:
    return {"update_type": "bot_started", "timestamp": 1, "chat_id": DIALOG, "user": user_obj(), "user_locale": "ru"}


def message_created(text: str = "", *, user_id: int = CLIENT, chat_id: int = DIALOG, chat_type: str = "dialog",
                    attachments: list[dict[str, Any]] | None = None, is_bot: bool = False) -> dict[str, Any]:
    sender = {**user_obj(user_id), "is_bot": is_bot}
    return {
        "update_type": "message_created",
        "timestamp": 1,
        "message": {
            "sender": sender,
            "recipient": {"chat_id": chat_id, "chat_type": chat_type},
            "timestamp": 1,
            "body": {"mid": "mid.in", "seq": 1, "text": text, "attachments": attachments or []},
        },
    }


def message_callback(payload: str, *, mid: str) -> dict[str, Any]:
    return {
        "update_type": "message_callback",
        "timestamp": 1,
        "callback": {"timestamp": 1, "callback_id": f"cb-{payload}", "payload": payload, "user": user_obj()},
        "message": {"recipient": {"chat_id": DIALOG, "chat_type": "dialog"}, "body": {"mid": mid, "seq": 1, "text": "..."}},
    }


VCARD = "BEGIN:VCARD\r\nVERSION:3.0\r\nFN:Мария\r\nTEL;TYPE=cell:+79001234567\r\nEND:VCARD\r\n"


@pytest.fixture
def api() -> FakeMaxApi:
    return FakeMaxApi()


@pytest.fixture
def transport(api: FakeMaxApi, config: BotConfig, storage: Storage, clock: FakeClock) -> MaxTransport:
    service = LeadService(config, storage, clock=clock)
    transport = MaxTransport(api, service, AdminAccess(chat_id=ADMIN_CHAT), webhook_secret="webhook-secret")
    service.add_notifier(transport.notifier())  # type: ignore[arg-type]
    return transport


async def test_bot_started_sends_consent_keyboard(transport: MaxTransport, api: FakeMaxApi, config: BotConfig) -> None:
    await transport.process_update(bot_started())
    [sent] = api.sent
    assert sent["chat_id"] == DIALOG and "согласие" in sent["text"]
    [keyboard] = sent["attachments"]
    assert keyboard["type"] == "inline_keyboard"
    rows = keyboard["payload"]["buttons"]
    assert rows[0] == [{"type": "link", "text": config.texts.policy_button, "url": config.privacy.policy_url}]
    assert rows[1] == [
        {"type": "callback", "text": config.texts.consent_accept_button, "payload": Payload.CONSENT_YES},
        {"type": "callback", "text": config.texts.consent_decline_button, "payload": Payload.CONSENT_NO},
    ]


async def test_full_flow_with_contact(transport: MaxTransport, api: FakeMaxApi, storage: Storage) -> None:
    await transport.process_update(bot_started())
    consent_mid, consent_text = api.last_mid(), api.sent[-1]["text"]
    await transport.process_update(message_callback(Payload.CONSENT_YES, mid=consent_mid))
    # Нажатие подтверждено: сообщение переписано тем же текстом без кнопок.
    assert api.answers[0]["message"] == {"text": consent_text, "format": "html", "attachments": []}

    await transport.process_update(message_callback(Payload.SERVICE_PREFIX + "bathroom", mid=api.last_mid()))
    await transport.process_update(message_created("Мария"))
    phone_prompt = api.sent[-1]
    assert phone_prompt["attachments"][0]["payload"]["buttons"] == [[{"type": "request_contact", "text": "Отправить мой номер"}]]

    contact = {"type": "contact", "payload": {"vcf_info": VCARD, "max_info": user_obj(), "hash": "x"}}
    await transport.process_update(message_created("", attachments=[contact]))
    assert "+7 900 123-45-67" in api.sent[-2]["text"]

    await transport.process_update(message_created("Удобно утром"))
    await transport.process_update(message_callback(Payload.SEND, mid=api.last_mid()))

    [lead] = storage.list_leads()
    assert (lead.channel, lead.user_id, lead.phone, lead.comment) == (Channel.MAX, CLIENT, "+79001234567", "Удобно утром")
    admin, thanks = api.sent[-2:]
    assert admin["chat_id"] == ADMIN_CHAT and "Новая заявка №1" in admin["text"] and "max://user/2002" in admin["text"]
    assert thanks["chat_id"] == DIALOG and "№1" in thanks["text"]


async def test_unknown_callback_message_gets_notification(transport: MaxTransport, api: FakeMaxApi, config: BotConfig) -> None:
    await transport.process_update(bot_started())
    await transport.process_update(message_callback(Payload.CONSENT_YES, mid="mid.unknown"))
    assert api.answers[0] == {"callback_id": "cb-consent:yes", "notification": config.texts.callback_ack, "message": None}


async def test_ignores_bots_and_group_chatter(transport: MaxTransport, api: FakeMaxApi) -> None:
    await transport.process_update(message_created("/start", is_bot=True))
    await transport.process_update(message_created("Хочу ремонт", chat_id=-1, chat_type="chat"))
    await transport.process_update({"update_type": "message_removed", "timestamp": 1})
    assert api.sent == []


async def test_admin_stats_and_export(transport: MaxTransport, api: FakeMaxApi, storage: Storage) -> None:
    await transport.process_update(message_created("/stats", user_id=9, chat_id=ADMIN_CHAT, chat_type="chat"))
    assert api.sent[-1]["chat_id"] == ADMIN_CHAT and "Всего: <b>0</b>" in api.sent[-1]["text"]

    await transport.process_update(message_created("/export", user_id=9, chat_id=ADMIN_CHAT, chat_type="chat"))
    assert api.sent[-1]["text"] == "Заявок пока нет." and api.uploads == []

    await transport.process_update(bot_started())
    await transport.process_update(message_callback(Payload.CONSENT_YES, mid=api.last_mid()))
    await transport.process_update(message_callback(Payload.SERVICE_PREFIX + "other", mid=api.last_mid()))
    await transport.process_update(message_created("Мария"))
    await transport.process_update(message_created("89001234567"))
    await transport.process_update(message_callback(Payload.SKIP_COMMENT, mid=api.last_mid()))
    await transport.process_update(message_callback(Payload.SEND, mid=api.last_mid()))

    await transport.process_update(message_created("/export", user_id=9, chat_id=ADMIN_CHAT, chat_type="chat"))
    [(filename, content)] = api.uploads
    assert filename.endswith(".xlsx")
    assert load_workbook(io.BytesIO(content))["Заявки"].cell(row=2, column=6).value == "+79001234567"
    assert api.sent[-1]["attachments"] == [{"type": "file", "payload": {"token": "file-token-1"}}]


async def test_export_upload_failure_is_reported(transport: MaxTransport, api: FakeMaxApi, storage: Storage) -> None:
    api.fail_upload = True
    await transport.process_update(bot_started())
    await transport.process_update(message_callback(Payload.CONSENT_YES, mid=api.last_mid()))
    await transport.process_update(message_callback(Payload.SERVICE_PREFIX + "other", mid=api.last_mid()))
    await transport.process_update(message_created("Мария"))
    await transport.process_update(message_created("89001234567"))
    await transport.process_update(message_callback(Payload.SKIP_COMMENT, mid=api.last_mid()))
    await transport.process_update(message_callback(Payload.SEND, mid=api.last_mid()))
    await transport.process_update(message_created("/export", user_id=9, chat_id=ADMIN_CHAT, chat_type="chat"))
    assert api.sent[-1]["text"] == ADMIN_EXPORT_FAILED


async def test_non_admin_cannot_get_stats(transport: MaxTransport, api: FakeMaxApi) -> None:
    await transport.process_update(message_created("/stats"))
    [reply] = api.sent
    assert "Всего" not in reply["text"] and "/start" in reply["text"]


async def test_polling_processes_batches_and_saves_marker(transport: MaxTransport, api: FakeMaxApi, storage: Storage) -> None:
    storage.set_value(MARKER_KEY, "5")
    api.batches = [
        ([bot_started()], 10),
        MaxApiError(502, None, "bad gateway"),  # временный сбой — повтор
        ([message_created("/cancel")], 11),
    ]
    await transport.run_polling(first_backoff=0, stop=api.stop)
    assert api.polled_markers[:4] == [5, 10, 10, 11]
    assert storage.get_value(MARKER_KEY) == "11"
    assert len(api.sent) == 2  # согласие и «заявка отменена»
    assert api.commands and api.commands[0][0] == "start"


async def test_polling_stops_on_bad_token(transport: MaxTransport, api: FakeMaxApi) -> None:
    api.batches = [MaxAuthError(401, "verify.token", "Invalid access_token")]
    with pytest.raises(MaxAuthError):
        await transport.run_polling(first_backoff=0, stop=api.stop)


async def test_broken_update_does_not_stop_polling(transport: MaxTransport, api: FakeMaxApi) -> None:
    api.batches = [([{"update_type": "bot_started", "user": {"no_user_id": True}}, bot_started()], 1)]
    await transport.run_polling(first_backoff=0, stop=api.stop)
    assert len(api.sent) == 1


async def test_webhook_checks_secret(transport: MaxTransport, api: FakeMaxApi) -> None:
    body = json.dumps(bot_started()).encode()
    assert await transport.handle_webhook(body, "wrong") == 403
    assert await transport.handle_webhook(body, None) == 403
    assert api.sent == []
    assert await transport.handle_webhook(b"not json", "webhook-secret") == 400
    assert await transport.handle_webhook(body, "webhook-secret") == 200
    assert len(api.sent) == 1
    app = transport.webhook_app("/max/webhook")
    assert {r.resource.canonical for r in app.router.routes()} >= {"/max/webhook", "/healthz"}  # type: ignore[union-attr]


@pytest.mark.parametrize(
    ("vcf", "expected"),
    [
        (VCARD, "+79001234567"),
        ("BEGIN:VCARD\nVERSION:4.0\nTEL;VALUE=uri;TYPE=cell:tel:+7-900-123-45-67\nEND:VCARD", "+7-900-123-45-67"),
        ("BEGIN:VCARD\nitem1.TEL:89001234567\nEND:VCARD", "89001234567"),
        ("BEGIN:VCARD\nFN:Без номера\nEND:VCARD", None),
        ("", None),
    ],
)
def test_phone_from_vcf(vcf: str, expected: str | None) -> None:
    assert phone_from_vcf(vcf) == expected


def test_contact_phone() -> None:
    assert contact_phone([{"type": "image", "payload": {}}]) is None
    assert contact_phone([{"type": "contact", "payload": {"vcf_info": "BEGIN:VCARD\nEND:VCARD"}}]) == ""
