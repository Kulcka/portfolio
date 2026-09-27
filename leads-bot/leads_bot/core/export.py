"""Выгрузка заявок в Excel (.xlsx)."""

from __future__ import annotations

import io
from collections import Counter
from collections.abc import Sequence
from datetime import datetime
from zoneinfo import ZoneInfo

from openpyxl import Workbook
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE, Cell
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from .models import Lead

DATETIME_FORMAT = "DD.MM.YYYY HH:MM"

# (заголовок, ширина колонки)
COLUMNS: tuple[tuple[str, int], ...] = (
    ("№", 7),
    ("Дата и время", 17),
    ("Канал", 10),
    ("Услуга", 24),
    ("Имя", 20),
    ("Телефон", 16),
    ("Комментарий", 50),
    ("Профиль в мессенджере", 24),
    ("ID пользователя", 16),
    ("Согласие на обработку ПДн", 19),
    ("Политика (версия)", 30),
    ("Уведомление отправлено", 14),
)

_HEADER_FONT = Font(bold=True, color="FFFFFF")
_HEADER_FILL = PatternFill("solid", fgColor="2F5597")


def _set_text(cell: Cell, value: str) -> None:
    """Записать строку как текст: без управляющих символов и без превращения `=...` в формулу."""
    cell.value = ILLEGAL_CHARACTERS_RE.sub("", value)
    if cell.data_type == "f":  # openpyxl считает строку с '=' формулой — клиентский ввод формулой не бывает
        cell.data_type = "s"


def _local(moment: datetime, tz: ZoneInfo) -> datetime:
    # Excel не хранит часовой пояс: пишем местное время без tzinfo.
    return moment.astimezone(tz).replace(tzinfo=None)


def build_workbook(leads: Sequence[Lead], tz: ZoneInfo) -> Workbook:
    wb = Workbook()
    ws = wb.active
    ws.title = "Заявки"
    ws.append([title for title, _ in COLUMNS])
    for index, (_, width) in enumerate(COLUMNS, start=1):
        ws.column_dimensions[get_column_letter(index)].width = width
        header = ws.cell(row=1, column=index)
        header.font = _HEADER_FONT
        header.fill = _HEADER_FILL
        header.alignment = Alignment(vertical="center", wrap_text=True)

    for row_index, lead in enumerate(leads, start=2):
        profile = f"@{lead.username}" if lead.username else (lead.display_name or "")
        policy = lead.policy_url + (f" ({lead.policy_version})" if lead.policy_version else "")
        ws.cell(row=row_index, column=1, value=lead.id)
        created = ws.cell(row=row_index, column=2, value=_local(lead.created_at, tz))
        created.number_format = DATETIME_FORMAT
        _set_text(ws.cell(row=row_index, column=3), lead.channel.title)
        _set_text(ws.cell(row=row_index, column=4), lead.service_title)
        _set_text(ws.cell(row=row_index, column=5), lead.name)
        _set_text(ws.cell(row=row_index, column=6), lead.phone)
        comment = ws.cell(row=row_index, column=7)
        _set_text(comment, lead.comment)
        comment.alignment = Alignment(wrap_text=True, vertical="top")
        _set_text(ws.cell(row=row_index, column=8), profile)
        _set_text(ws.cell(row=row_index, column=9), str(lead.user_id))
        consent = ws.cell(row=row_index, column=10, value=_local(lead.consent_at, tz))
        consent.number_format = DATETIME_FORMAT
        _set_text(ws.cell(row=row_index, column=11), policy)
        _set_text(ws.cell(row=row_index, column=12), "да" if lead.admin_notified else "нет")

    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(COLUMNS))}{max(len(leads) + 1, 2)}"

    summary = wb.create_sheet("Сводка")
    summary.append(["Услуга", "Заявок"])
    for title, count in Counter(lead.service_title for lead in leads).most_common():
        summary.append([None, count])
        _set_text(summary.cell(row=summary.max_row, column=1), title)
    summary.append([])
    summary.append(["Всего", len(leads)])
    summary.column_dimensions["A"].width = 30
    summary.column_dimensions["B"].width = 10
    for cell in summary[1]:
        cell.font = _HEADER_FONT
        cell.fill = _HEADER_FILL
    summary.cell(row=summary.max_row, column=1).font = Font(bold=True)
    return wb


def export_leads_xlsx(leads: Sequence[Lead], tz: ZoneInfo) -> bytes:
    buffer = io.BytesIO()
    build_workbook(leads, tz).save(buffer)
    return buffer.getvalue()


def export_filename(now: datetime, tz: ZoneInfo) -> str:
    return f"leads_{now.astimezone(tz).strftime('%Y-%m-%d_%H%M')}.xlsx"
